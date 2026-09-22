"""Abnahmetests für die Durchsicht der Rollkandidaten (AP-04d).

Geschrieben nach §11.2 des Plans durch das starke Modell (AP-04dv) — für das
günstige Modell unveränderlich. Maßgeblich ist der Vertrag in
`auftraege/AP-04dv.md`, fachlich Plan §6, AP-04, Block B, Regel 2
(„Durchsicht über eine Datei-Rundreise, bis AP-08 steht").

Das Modul springt still ab, solange `journal/durchsicht.py` fehlt.

Alle Szenarien sind frei erfunden (`tests/fixtures_rebuild.py`,
`tests/fixtures_durchsicht.py`). `fixtures/golden/` wird nicht gelesen.
**Die Preise sind von Hand nachgerechnet** (Regel 8) — siehe
`test_preise_sind_von_hand_nachgerechnet`.
"""

from datetime import datetime
from pathlib import Path

import pytest
from fixtures_durchsicht import DOPPELT, INSTRUMENTE_ESTX, MEHRDEUTIG
from fixtures_rebuild import befuelle
from openpyxl import load_workbook
from sqlalchemy import select

pytest.importorskip("journal.durchsicht")

from journal.cli import main  # noqa: E402
from journal.db import create_db_engine  # noqa: E402
from journal.durchsicht import kandidaten_ausschreiben, urteile_einlesen  # noqa: E402
from journal.rebuild import rebuild  # noqa: E402
from journal.schema import chain_links, chain_overrides, chains, structures  # noqa: E402

PROJEKTWURZEL = Path(__file__).resolve().parents[1]
ZEITPUNKT = datetime(2026, 9, 22, 9, 0, 0)

KOPF = [
    "Gruppe", "Kandidaten in Gruppe", "Rang", "Basiswert", "Seite",
    "Alt: Legs", "Alt: geschlossen", "Alt: Preis zu",
    "Neu: Legs", "Neu: eröffnet", "Neu: Preis auf",
    "Abstand Minuten", "Mengenverhältnis", "Entscheidung", "Bemerkung",
    "Schlüssel Kette", "Schlüssel Struktur",
]


# ---------------------------------------------------------------------------
# Hilfen
# ---------------------------------------------------------------------------


def _bestand(db_pfad, ausfuehrungen=None, instrumente=None):
    befuelle(db_pfad, ausfuehrungen=ausfuehrungen, instrumente_=instrumente)
    rebuild(db_pfad)


def _zeilen(datei) -> list[dict]:
    blatt = load_workbook(datei)["Kandidaten"]
    werte = list(blatt.iter_rows(values_only=True))
    kopf = list(werte[0])
    return [dict(zip(kopf, zeile, strict=True)) for zeile in werte[1:] if any(zeile)]


def _entscheide(datei, entscheidungen: dict, bemerkungen: dict | None = None) -> None:
    """Setzt `Entscheidung` (und `Bemerkung`) je (Gruppe, Rang) — so, wie der
    Nutzer es in der Tabellenkalkulation täte."""
    mappe = load_workbook(datei)
    blatt = mappe["Kandidaten"]
    kopf = [zelle.value for zelle in blatt[1]]
    spalte_e = kopf.index("Entscheidung") + 1
    spalte_b = kopf.index("Bemerkung") + 1
    spalte_g = kopf.index("Gruppe")
    spalte_r = kopf.index("Rang")
    for zeile in range(2, blatt.max_row + 1):
        werte = [zelle.value for zelle in blatt[zeile]]
        schluessel = (werte[spalte_g], werte[spalte_r])
        if schluessel in entscheidungen:
            blatt.cell(zeile, spalte_e).value = entscheidungen[schluessel]
        if bemerkungen and schluessel in bemerkungen:
            blatt.cell(zeile, spalte_b).value = bemerkungen[schluessel]
    mappe.save(datei)


def _overrides(db_pfad) -> list[dict]:
    engine = create_db_engine(db_pfad)
    with engine.connect() as verbindung:
        return [dict(z) for z in verbindung.execute(select(chain_overrides)).mappings()]


def _schluessel_je_id(db_pfad, tabelle, spalte) -> dict:
    engine = create_db_engine(db_pfad)
    with engine.connect() as verbindung:
        return {
            z[spalte]: z["natur_schluessel"]
            for z in verbindung.execute(select(tabelle)).mappings()
        }


def _urteile(db_pfad) -> dict:
    """chain_overrides als {(Kettenschlüssel, Strukturschlüssel): urteil}."""
    ketten = _schluessel_je_id(db_pfad, chains, "chain_id")
    strukturen = _schluessel_je_id(db_pfad, structures, "structure_id")
    return {
        (ketten[o["chain_id"]], strukturen[o["structure_id"]]): o["urteil"]
        for o in _overrides(db_pfad)
    }


# ---------------------------------------------------------------------------
# Ausschreiben
# ---------------------------------------------------------------------------


def test_ausschreiben_schreibt_offene_kandidaten_mit_kopf(migrated_db, tmp_path):
    """Vertrag 3–5: Blatt `Kandidaten`, Kopf in Zeile 1 in fester Reihenfolge,
    eine Zeile je offenem Kandidaten. Der Standardbestand hat genau einen:
    den ESTX50-Roll in zwei Orders."""
    _bestand(migrated_db)
    datei = tmp_path / "kandidaten.xlsx"

    ergebnis = kandidaten_ausschreiben(migrated_db, datei)

    assert ergebnis.gruppen == 1
    assert ergebnis.kandidaten == 1
    blatt = load_workbook(datei)["Kandidaten"]
    assert [zelle.value for zelle in blatt[1]] == KOPF
    zeilen = _zeilen(datei)
    assert len(zeilen) == 1
    zeile = zeilen[0]
    assert zeile["Gruppe"] == 1
    assert zeile["Kandidaten in Gruppe"] == 1
    assert zeile["Rang"] == 1
    assert zeile["Basiswert"] == "ESTX50"
    assert zeile["Seite"] == "Put"
    assert zeile["Schlüssel Kette"] == "k:Z1|Z2"
    assert zeile["Schlüssel Struktur"] == "Z5|Z6"
    assert zeile["Entscheidung"] is None
    assert "4900" in zeile["Alt: Legs"] and "4850" in zeile["Alt: Legs"]
    assert "4800" in zeile["Neu: Legs"] and "4750" in zeile["Neu: Legs"]
    assert zeile["Alt: geschlossen"] == datetime(2026, 9, 16, 15, 0, 0)
    assert zeile["Neu: eröffnet"] == datetime(2026, 9, 16, 15, 2, 0)


def test_preise_sind_von_hand_nachgerechnet(migrated_db, tmp_path):
    """Vertrag 6: Nettoprämie je Kontrakt, Verkauf positiv, Kauf negativ.

    Von Hand, Menge 1:
    - Alt, geschlossen: Z3 kauft 4900 zurück zu 12,00 (−12), Z4 verkauft
      4850 zu 6,00 (+6) → **−6,00**
    - Neu, eröffnet: Z5 verkauft 4800 zu 28,00 (+28), Z6 kauft 4750 zu
      17,00 (−17) → **+11,00**
    - Abstand 15:00 → 15:02 = **2** Minuten, Mengenverhältnis **1**.
    """
    _bestand(migrated_db)
    datei = tmp_path / "kandidaten.xlsx"
    kandidaten_ausschreiben(migrated_db, datei)

    zeile = _zeilen(datei)[0]

    assert zeile["Alt: Preis zu"] == pytest.approx(-6.00)
    assert zeile["Neu: Preis auf"] == pytest.approx(11.00)
    assert zeile["Abstand Minuten"] == 2
    assert zeile["Mengenverhältnis"] == 1


def test_mehrdeutige_gruppe_zeigt_alle_fortsetzungen(migrated_db, tmp_path):
    """Vertrag 4: Eine schließende Position mit zwei möglichen Fortsetzungen
    ist **eine** Gruppe mit zwei Zeilen, nach Rang geordnet. Genau diese
    Fälle soll der Nutzer vergleichen können (Plan §6, B/2)."""
    _bestand(migrated_db, MEHRDEUTIG, INSTRUMENTE_ESTX)
    datei = tmp_path / "kandidaten.xlsx"

    ergebnis = kandidaten_ausschreiben(migrated_db, datei)

    assert ergebnis.gruppen == 1
    assert ergebnis.kandidaten == 2
    zeilen = _zeilen(datei)
    assert [(z["Gruppe"], z["Rang"]) for z in zeilen] == [(1, 1), (1, 2)]
    assert all(z["Kandidaten in Gruppe"] == 2 for z in zeilen)
    assert [z["Schlüssel Struktur"] for z in zeilen] == ["Y5|Y6", "Y7|Y8"]
    assert [z["Abstand Minuten"] for z in zeilen] == [2, 20]


def test_entschiedene_kandidaten_werden_nicht_ausgeschrieben(migrated_db, tmp_path):
    """Vertrag 2: Nur **offene** Kandidaten. Nach einer Ablehnung taucht der
    Kandidat nicht wieder auf — das ist der Zweck gespeicherter
    Ablehnungen (Plan §6, B/2)."""
    _bestand(migrated_db)
    erste = tmp_path / "erste.xlsx"
    kandidaten_ausschreiben(migrated_db, erste)
    _entscheide(erste, {(1, 1): "nein"})
    urteile_einlesen(migrated_db, erste, jetzt=ZEITPUNKT)
    rebuild(migrated_db)

    zweite = tmp_path / "zweite.xlsx"
    ergebnis = kandidaten_ausschreiben(migrated_db, zweite)

    assert ergebnis.kandidaten == 0
    assert _zeilen(zweite) == []


def test_ausschreiben_ueberschreibt_keine_datei(migrated_db, tmp_path):
    """Vertrag 7: Eine vorhandene Datei kann Entscheidungen tragen."""
    _bestand(migrated_db)
    datei = tmp_path / "kandidaten.xlsx"
    datei.write_bytes(b"vorhanden")

    with pytest.raises(FileExistsError):
        kandidaten_ausschreiben(migrated_db, datei)
    assert datei.read_bytes() == b"vorhanden"


def test_ausschreiben_verweigert_ziel_im_repository(migrated_db):
    """Vertrag 8, Regel 2: Die Datei trägt echte Kontodaten, sobald sie aus
    der echten Datenbank kommt. Im Repository — außer unter
    `fixtures/golden/`, das nicht versioniert wird — darf sie nicht landen."""
    _bestand(migrated_db)
    ziel = PROJEKTWURZEL / "kandidaten_pruefung_ap04dv.xlsx"

    with pytest.raises(ValueError):
        kandidaten_ausschreiben(migrated_db, ziel)
    assert not ziel.exists()


def test_ausschreiben_verlangt_vorherigen_rebuild(migrated_db, tmp_path):
    """Vertrag 9: Die Schlüssel in der Datei müssen auf Zeilen in
    `structures` und `chains` zeigen. Ohne `journal rebuild` gibt es die
    nicht — dann benannter Abbruch statt einer Datei, die sich nie
    zurücklesen ließe."""
    befuelle(migrated_db)
    datei = tmp_path / "kandidaten.xlsx"

    with pytest.raises(ValueError, match="rebuild"):
        kandidaten_ausschreiben(migrated_db, datei)
    assert not datei.exists()


# ---------------------------------------------------------------------------
# Einlesen
# ---------------------------------------------------------------------------


def test_einlesen_ja_und_nein(migrated_db, tmp_path):
    """Vertrag 11: `ja` wird Bestätigung, `nein` Ablehnung, leer bleibt
    offen. Beide Urteile landen in `chain_overrides`."""
    _bestand(migrated_db, MEHRDEUTIG, INSTRUMENTE_ESTX)
    datei = tmp_path / "kandidaten.xlsx"
    kandidaten_ausschreiben(migrated_db, datei)
    _entscheide(datei, {(1, 1): "ja", (1, 2): "Nein"})

    ergebnis = urteile_einlesen(migrated_db, datei, jetzt=ZEITPUNKT)

    assert ergebnis.neu_bestaetigt == 1
    assert ergebnis.neu_abgelehnt == 1
    assert ergebnis.unentschieden == 0
    assert _urteile(migrated_db) == {
        ("k:Y1|Y2", "Y5|Y6"): "bestaetigt",
        ("k:Y1|Y2", "Y7|Y8"): "abgelehnt",
    }
    assert all(o["erstellt_am"] == ZEITPUNKT for o in _overrides(migrated_db))


def test_leere_entscheidung_bleibt_offen(migrated_db, tmp_path):
    """Vertrag 11: Wer eine Zeile nicht entscheidet, entscheidet nichts."""
    _bestand(migrated_db, MEHRDEUTIG, INSTRUMENTE_ESTX)
    datei = tmp_path / "kandidaten.xlsx"
    kandidaten_ausschreiben(migrated_db, datei)
    _entscheide(datei, {(1, 1): "ja"})

    ergebnis = urteile_einlesen(migrated_db, datei, jetzt=ZEITPUNKT)

    assert ergebnis.neu_bestaetigt == 1
    assert ergebnis.unentschieden == 1
    assert len(_overrides(migrated_db)) == 1


def test_einlesen_ist_idempotent(migrated_db, tmp_path):
    """Vertrag 13: Dieselbe Datei zweimal eingelesen legt nichts doppelt an."""
    _bestand(migrated_db, MEHRDEUTIG, INSTRUMENTE_ESTX)
    datei = tmp_path / "kandidaten.xlsx"
    kandidaten_ausschreiben(migrated_db, datei)
    _entscheide(datei, {(1, 1): "ja", (1, 2): "nein"})
    urteile_einlesen(migrated_db, datei, jetzt=ZEITPUNKT)
    vorher = _overrides(migrated_db)

    zweites = urteile_einlesen(migrated_db, datei, jetzt=datetime(2026, 9, 23, 9, 0, 0))

    assert zweites.neu_bestaetigt == 0
    assert zweites.neu_abgelehnt == 0
    assert zweites.unveraendert == 2
    assert _overrides(migrated_db) == vorher


def test_widerspruch_zu_bestehender_entscheidung_bricht_ab(migrated_db, tmp_path):
    """Vertrag 14, Regel 4: Eine getroffene Entscheidung wird nicht still
    überschrieben. Widerspricht die Datei einer gespeicherten, bricht das
    Einlesen ab — und schreibt **nichts**, auch nicht die übrigen Zeilen."""
    _bestand(migrated_db, MEHRDEUTIG, INSTRUMENTE_ESTX)
    erste = tmp_path / "erste.xlsx"
    kandidaten_ausschreiben(migrated_db, erste)
    _entscheide(erste, {(1, 1): "ja"})
    urteile_einlesen(migrated_db, erste, jetzt=ZEITPUNKT)
    vorher = _overrides(migrated_db)

    zweite = tmp_path / "zweite.xlsx"
    _entscheide(erste, {(1, 1): "nein", (1, 2): "nein"})
    zweite.write_bytes(erste.read_bytes())

    with pytest.raises(ValueError):
        urteile_einlesen(migrated_db, zweite, jetzt=ZEITPUNKT)
    assert _overrides(migrated_db) == vorher


def test_ueberschreiben_aendert_die_entscheidung(migrated_db, tmp_path):
    """Vertrag 14: Wer seine Meinung ändert, sagt es ausdrücklich."""
    _bestand(migrated_db, MEHRDEUTIG, INSTRUMENTE_ESTX)
    datei = tmp_path / "kandidaten.xlsx"
    kandidaten_ausschreiben(migrated_db, datei)
    _entscheide(datei, {(1, 1): "ja"})
    urteile_einlesen(migrated_db, datei, jetzt=ZEITPUNKT)

    _entscheide(datei, {(1, 1): "nein"})
    ergebnis = urteile_einlesen(migrated_db, datei, ueberschreiben=True, jetzt=ZEITPUNKT)

    assert ergebnis.geaendert == 1
    assert _urteile(migrated_db) == {("k:Y1|Y2", "Y5|Y6"): "abgelehnt"}


def test_ungueltige_entscheidung_bricht_ab_und_schreibt_nichts(migrated_db, tmp_path):
    """Vertrag 12: Nur `ja` und `nein` (auch `j`/`n`, Groß-/Kleinschreibung
    egal). Alles andere ist ein Abbruch mit Nennung von Gruppe und Rang —
    und die gültige Zeile daneben wird ebenfalls nicht geschrieben
    (Vertrag 15: alles oder nichts)."""
    _bestand(migrated_db, MEHRDEUTIG, INSTRUMENTE_ESTX)
    datei = tmp_path / "kandidaten.xlsx"
    kandidaten_ausschreiben(migrated_db, datei)
    _entscheide(datei, {(1, 1): "ja", (1, 2): "vielleicht"})

    with pytest.raises(ValueError, match="Gruppe 1"):
        urteile_einlesen(migrated_db, datei, jetzt=ZEITPUNKT)
    assert _overrides(migrated_db) == []


def test_wahrheitswert_aus_excel_wird_abgelehnt(migrated_db, tmp_path):
    """Vertrag 12: Tippt man in Excel „Falsch" oder „Wahr", wird daraus ein
    Wahrheitswert, kein Text. Er ist keine gültige Entscheidung — sonst würde
    ein Tippfehler still als Urteil gelesen (Befund aus der Stichprobe,
    OF-41)."""
    _bestand(migrated_db)
    datei = tmp_path / "kandidaten.xlsx"
    kandidaten_ausschreiben(migrated_db, datei)
    _entscheide(datei, {(1, 1): False})

    with pytest.raises(ValueError):
        urteile_einlesen(migrated_db, datei, jetzt=ZEITPUNKT)
    assert _overrides(migrated_db) == []


def test_doppelte_beanspruchung_bricht_ab(migrated_db, tmp_path):
    """Vertrag 16 (AP-04v, Vertrag 21): Zwei Positionen schließen, eine neue
    öffnet und ist Kandidat für beide. Bestätigt der Nutzer beide, gehörte
    dieselbe Struktur in zwei Ketten — Abbruch, nichts geschrieben."""
    _bestand(migrated_db, DOPPELT, INSTRUMENTE_ESTX)
    datei = tmp_path / "kandidaten.xlsx"
    ergebnis = kandidaten_ausschreiben(migrated_db, datei)
    assert ergebnis.gruppen == 2
    gruppen = {z["Gruppe"] for z in _zeilen(datei)}
    _entscheide(datei, {(g, 1): "ja" for g in gruppen})

    with pytest.raises(ValueError):
        urteile_einlesen(migrated_db, datei, jetzt=ZEITPUNKT)
    assert _overrides(migrated_db) == []


def test_bemerkung_wird_zur_notiz(migrated_db, tmp_path):
    """Vertrag 11: Die Bemerkung des Nutzers bleibt als `notiz` erhalten."""
    _bestand(migrated_db)
    datei = tmp_path / "kandidaten.xlsx"
    kandidaten_ausschreiben(migrated_db, datei)
    _entscheide(datei, {(1, 1): "nein"}, {(1, 1): "Re-Entry nach PT"})

    urteile_einlesen(migrated_db, datei, jetzt=ZEITPUNKT)

    assert _overrides(migrated_db)[0]["notiz"] == "Re-Entry nach PT"


# ---------------------------------------------------------------------------
# Rundreise
# ---------------------------------------------------------------------------


def test_rundreise_ende_zu_ende(migrated_db, tmp_path):
    """Ausschreiben → entscheiden → einlesen → `rebuild`: Der bestätigte
    Kandidat ist danach ein Kettenglied mit `erkennung = 'manuell'`, und es
    ist kein Kandidat mehr offen."""
    _bestand(migrated_db)
    datei = tmp_path / "kandidaten.xlsx"
    kandidaten_ausschreiben(migrated_db, datei)
    _entscheide(datei, {(1, 1): "ja"})
    urteile_einlesen(migrated_db, datei, jetzt=ZEITPUNKT)

    ergebnis = rebuild(migrated_db)

    assert ergebnis.kandidaten_offen == 0
    assert ergebnis.overrides_angewandt == 1
    ketten = _schluessel_je_id(migrated_db, chains, "chain_id")
    strukturen = _schluessel_je_id(migrated_db, structures, "structure_id")
    engine = create_db_engine(migrated_db)
    with engine.connect() as verbindung:
        glieder = {
            strukturen[z["structure_id"]]: (ketten[z["chain_id"]], z["erkennung"])
            for z in verbindung.execute(select(chain_links)).mappings()
        }
    assert glieder["Z5|Z6"] == ("k:Z1|Z2", "manuell")


def test_cli_rundreise(migrated_db, tmp_path, capsys):
    """Vertrag 17: beide Kommandos über `journal`, Rückgabe 0 im Erfolgsfall,
    1 bei einem Fehler."""
    _bestand(migrated_db)
    datei = tmp_path / "kandidaten.xlsx"

    assert main(["kandidaten-ausschreiben", str(datei), "--db", str(migrated_db)]) == 0
    _entscheide(datei, {(1, 1): "ja"})
    assert main(["kandidaten-einlesen", str(datei), "--db", str(migrated_db)]) == 0
    assert "rebuild" in capsys.readouterr().out

    assert main(["kandidaten-ausschreiben", str(datei), "--db", str(migrated_db)]) == 1
