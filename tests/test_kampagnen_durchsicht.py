"""Abnahmetests für AP-05d: Klassifikation der Campaign-Kandidaten über eine
Datei-Rundreise (Ausschreiben, Entscheiden, Einlesen) und ihre Persistenz.

Geschrieben nach §11.2 des Plans durch das starke Modell — für das
günstige Modell unveränderlich. Maßgeblich ist der Vertrag in
`auftraege/AP-05d.md`; die Entscheidungen stehen in `OFFENE_FRAGEN.md`,
OF-54 und OF-71.

Testwelt (frei erfunden, DAX, alles Kontosicht, Multiplikator 5):

  P   Put-Vertikale, eröffnet 01.09. 10:00, am 15.09. 15:00 in einer Order
      in den Oktober gerollt (Fixture `fixtures_rebuild`, Ausführungen E1-E6),
      offen. Naturschlüssel der Kette `k:E1|E2`, zwei Glieder, ein Roll.
  C1  Call-Vertikale, Expiry 18.09., eröffnet 02.09. 10:00, offen
      (F1/F2, Kette `k:F1|F2`).
  C2  Call-Vertikale, Expiry 20.11., eröffnet 15.09. 17:00 - zwei Stunden nach
      dem Roll, offen (G1/G2, Kette `k:G1|G2`).

Erwartete Kandidaten (Kette A ist die früher eröffnete):
  financing  (P, C2)  Abstand 120 Minuten, beide offen -> Überlappung unbegrenzt
  ic         (P, C1)  P's erstes Glied (Expiry 18.09.) und C1 überlappen;
                      Abstand 24 Stunden = 1440 Minuten, beide offen.
  Kein IC (P, C2): C2 trägt die Expiry 20.11., P nur 18.09. und 16.10.
"""

from datetime import date, datetime
from decimal import Decimal

import pytest
from fixtures_rebuild import AUSFUEHRUNGEN as FIXTURE_AUSFUEHRUNGEN
from fixtures_rebuild import INSTRUMENTE as FIXTURE_INSTRUMENTE
from fixtures_rebuild import befuelle
from openpyxl import load_workbook
from sqlalchemy import delete, insert, select, update

from journal.cli import main
from journal.db import create_db_engine
from journal.rebuild import rebuild
from journal.schema import (
    campaign_ablehnungen,
    campaign_links,
    campaigns,
    chains,
    executions,
    legs,
)

try:
    from journal.kampagnen_durchsicht import (
        kampagnen_ausschreiben,
        kampagnen_urteile_einlesen,
    )
except ModuleNotFoundError:  # Startzustand: ROT, siehe Vertrag
    kampagnen_ausschreiben = kampagnen_urteile_einlesen = None

KOPF = [
    "Gruppe", "Kandidaten in Gruppe", "Rang", "Typ", "Basiswert",
    "Kette A: Seite", "Kette A: eröffnet", "Kette A: beendet", "Kette A: Rolls",
    "Kette B: Seite", "Kette B: eröffnet", "Kette B: beendet", "Kette B: Rolls",
    "Abstand Minuten", "Überlappung Tage", "Typ Kette A", "Typ Kette B",
    "Entscheidung", "Bemerkung", "Schlüssel Kette A", "Schlüssel Kette B",
]

JETZT = datetime(2026, 10, 9, 9, 0, 0)

_INSTRUMENTE = [i for i in FIXTURE_INSTRUMENTE if i["conid"] in (601, 602, 701, 702)] + [
    dict(conid=611, symbol="DAX  260918C19500", underlying="DAX", asset_class="OPT",
         expiry=date(2026, 9, 18), strike=Decimal("19500"), right="C",
         multiplikator=Decimal("5"), waehrung="EUR"),
    dict(conid=612, symbol="DAX  260918C19600", underlying="DAX", asset_class="OPT",
         expiry=date(2026, 9, 18), strike=Decimal("19600"), right="C",
         multiplikator=Decimal("5"), waehrung="EUR"),
    dict(conid=621, symbol="DAX  261120C19500", underlying="DAX", asset_class="OPT",
         expiry=date(2026, 11, 20), strike=Decimal("19500"), right="C",
         multiplikator=Decimal("5"), waehrung="EUR"),
    dict(conid=622, symbol="DAX  261120C19600", underlying="DAX", asset_class="OPT",
         expiry=date(2026, 11, 20), strike=Decimal("19600"), right="C",
         multiplikator=Decimal("5"), waehrung="EUR"),
]

_AUSFUEHRUNGEN = [s for s in FIXTURE_AUSFUEHRUNGEN if s["exec_id"].startswith("E")] + [
    dict(exec_id="F1", conid=611, side="SELL", preis=Decimal("30.00"),
         handelszeit=datetime(2026, 9, 2, 10, 0, 0), exec_batch="5.1"),
    dict(exec_id="F2", conid=612, side="BUY", preis=Decimal("20.00"),
         handelszeit=datetime(2026, 9, 2, 10, 0, 0), exec_batch="5.1"),
    dict(exec_id="G1", conid=621, side="SELL", preis=Decimal("40.00"),
         handelszeit=datetime(2026, 9, 15, 17, 0, 0), exec_batch="6.1"),
    dict(exec_id="G2", conid=622, side="BUY", preis=Decimal("25.00"),
         handelszeit=datetime(2026, 9, 15, 17, 0, 0), exec_batch="6.1"),
]

P, C1, C2 = "k:E1|E2", "k:F1|F2", "k:G1|G2"


@pytest.fixture
def welt(migrated_db):
    befuelle(migrated_db, ausfuehrungen=_AUSFUEHRUNGEN, instrumente_=_INSTRUMENTE)
    rebuild(migrated_db)
    return migrated_db


def _ketten(db) -> dict:
    engine = create_db_engine(db)
    with engine.connect() as v:
        zeilen = v.execute(select(chains).where(chains.c.sicht == "konto")).mappings().all()
    return {z["natur_schluessel"]: dict(z) for z in zeilen}


def _alle(db, tabelle) -> list[dict]:
    engine = create_db_engine(db)
    with engine.connect() as v:
        return [dict(z) for z in v.execute(select(tabelle)).mappings().all()]


def _datei(welt, tmp_path, name="k.xlsx"):
    ziel = tmp_path / name
    kampagnen_ausschreiben(welt, ziel)
    return ziel


def _zeilen(datei) -> list[dict]:
    blatt = load_workbook(datei, data_only=True)["Kandidaten"]
    kopf = [c.value for c in blatt[1]]
    return [dict(zip(kopf, z, strict=True)) for z in blatt.iter_rows(min_row=2, values_only=True)]


def _entscheide(datei, entscheidungen: dict, **spalten):
    """`entscheidungen`: {(Typ, Schlüssel Kette B): Wert}; `spalten`:
    zusätzliche Zellwerte je Schlüssel, z. B. typ_a={(...): 'contract_roll'}."""
    mappe = load_workbook(datei)
    blatt = mappe["Kandidaten"]
    kopf = [c.value for c in blatt[1]]
    spalte = {name: kopf.index(name) + 1 for name in kopf}
    for zeile in range(2, blatt.max_row + 1):
        schluessel = (blatt.cell(zeile, spalte["Typ"]).value,
                      blatt.cell(zeile, spalte["Schlüssel Kette B"]).value)
        if schluessel in entscheidungen:
            blatt.cell(zeile, spalte["Entscheidung"], entscheidungen[schluessel])
        for name, werte in spalten.items():
            if schluessel in werte:
                blatt.cell(zeile, spalte[name], werte[schluessel])
    mappe.save(datei)


FIN = ("financing", C2)
IC = ("ic", C1)


# ---------------------------------------------------------------- Kandidaten


def test_welt_hat_genau_diese_beiden_offenen_kandidaten(migrated_db):
    befuelle(migrated_db, ausfuehrungen=_AUSFUEHRUNGEN, instrumente_=_INSTRUMENTE)

    ergebnis = rebuild(migrated_db)

    assert ergebnis.campaign_kandidaten_offen == 2


# ---------------------------------------------------------------- Ausschreiben


def test_ausschreiben_gruppiert_nach_kette_a_und_sortiert_financing_zuerst(welt, tmp_path):
    ziel = tmp_path / "k.xlsx"

    ergebnis = kampagnen_ausschreiben(welt, ziel)

    assert (ergebnis.gruppen, ergebnis.kandidaten) == (1, 2)
    mappe = load_workbook(ziel)
    assert mappe.sheetnames == ["Kandidaten", "Anleitung"]
    assert [c.value for c in mappe["Kandidaten"][1]] == KOPF
    financing, ic = _zeilen(ziel)
    assert (financing["Gruppe"], financing["Kandidaten in Gruppe"], financing["Rang"]) == (1, 2, 1)
    assert (ic["Gruppe"], ic["Kandidaten in Gruppe"], ic["Rang"]) == (1, 2, 2)


def test_ausschreiben_zeilen_financing(welt, tmp_path):
    financing = _zeilen(_datei(welt, tmp_path))[0]

    assert financing["Typ"] == "financing"
    assert financing["Basiswert"] == "DAX"
    assert (financing["Kette A: Seite"], financing["Kette A: eröffnet"],
            financing["Kette A: beendet"], financing["Kette A: Rolls"]) == (
        "Put", datetime(2026, 9, 1, 10, 0), None, 1)
    assert (financing["Kette B: Seite"], financing["Kette B: eröffnet"],
            financing["Kette B: beendet"], financing["Kette B: Rolls"]) == (
        "Call", datetime(2026, 9, 15, 17, 0), None, 0)
    assert financing["Abstand Minuten"] == 120
    assert financing["Überlappung Tage"] is None  # beide offen: unbegrenzt
    assert (financing["Typ Kette A"], financing["Typ Kette B"]) == ("ic_leg", "financing")
    assert financing["Entscheidung"] is None
    assert (financing["Schlüssel Kette A"], financing["Schlüssel Kette B"]) == (P, C2)


def test_ausschreiben_zeilen_ic(welt, tmp_path):
    ic = _zeilen(_datei(welt, tmp_path))[1]

    assert ic["Typ"] == "ic"
    assert ic["Kette B: eröffnet"] == datetime(2026, 9, 2, 10, 0)
    assert ic["Abstand Minuten"] == 1440
    assert ic["Überlappung Tage"] is None
    assert (ic["Typ Kette A"], ic["Typ Kette B"]) == ("ic_leg", "ic_leg")
    assert (ic["Schlüssel Kette A"], ic["Schlüssel Kette B"]) == (P, C1)


def test_ausschreiben_schuetzt_vorhandene_dateien_und_das_repository(welt, tmp_path):
    ziel = _datei(welt, tmp_path)

    with pytest.raises(FileExistsError):
        kampagnen_ausschreiben(welt, ziel)
    with pytest.raises(ValueError, match="Projektverzeichnis"):
        kampagnen_ausschreiben(welt, "kampagnen_test_nicht_im_repo.xlsx")


# ---------------------------------------------------------------- Einlesen: Bestätigung


def test_bestaetigter_ic_legt_eine_kampagne_an(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: "ja"})

    ergebnis = kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)

    assert (ergebnis.neu_bestaetigt, ergebnis.neu_abgelehnt, ergebnis.kampagnen_neu,
            ergebnis.verknuepfungen_neu, ergebnis.unentschieden) == (1, 0, 1, 2, 1)
    ketten = _ketten(welt)
    (kampagne,) = _alle(welt, campaigns)
    assert kampagne["underlying"] == "DAX"
    assert kampagne["eroeffnet_am"] == datetime(2026, 9, 1, 10, 0)  # früheste Eröffnung
    assert kampagne["beendet_am"] is None  # beide Ketten offen
    links = {z["chain_id"]: z for z in _alle(welt, campaign_links)}
    assert set(links) == {ketten[P]["chain_id"], ketten[C1]["chain_id"]}
    for link in links.values():
        assert link["campaign_id"] == kampagne["campaign_id"]
        assert link["beziehungstyp"] == "ic_leg"
        assert link["erkennung"] == "manuell"
        assert link["bestaetigt_am"] == JETZT


def test_typ_der_bestandskette_ist_aenderbar(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {FIN: "ja"}, **{"Typ Kette A": {FIN: "contract_roll"}})

    kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)

    ketten = _ketten(welt)
    typ = {z["chain_id"]: z["beziehungstyp"] for z in _alle(welt, campaign_links)}
    assert typ[ketten[P]["chain_id"]] == "contract_roll"
    assert typ[ketten[C2]["chain_id"]] == "financing"


def test_zwei_bestaetigungen_mit_gemeinsamer_kette_ergeben_eine_kampagne(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {FIN: "ja", IC: "ja"})

    ergebnis = kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)

    assert (ergebnis.neu_bestaetigt, ergebnis.kampagnen_neu, ergebnis.verknuepfungen_neu) == (
        2, 1, 3)
    assert len(_alle(welt, campaigns)) == 1
    typ = {z["chain_id"]: z["beziehungstyp"] for z in _alle(welt, campaign_links)}
    ketten = _ketten(welt)
    assert typ == {ketten[P]["chain_id"]: "ic_leg", ketten[C1]["chain_id"]: "ic_leg",
                   ketten[C2]["chain_id"]: "financing"}


def test_widerspruechliche_typen_derselben_kette_in_der_datei_brechen_ab(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {FIN: "ja", IC: "ja"}, **{"Typ Kette A": {FIN: "contract_roll"}})

    with pytest.raises(ValueError, match="Typ"):
        kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)

    assert _alle(welt, campaigns) == []
    assert _alle(welt, campaign_links) == []


def test_zweiter_lauf_derselben_datei_aendert_nichts(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: "ja"})
    kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)
    vorher = (_alle(welt, campaigns), _alle(welt, campaign_links))

    ergebnis = kampagnen_urteile_einlesen(welt, ziel, jetzt=datetime(2026, 10, 10))

    assert (ergebnis.neu_bestaetigt, ergebnis.unveraendert) == (0, 1)
    assert (_alle(welt, campaigns), _alle(welt, campaign_links)) == vorher


def test_zwei_verschiedene_kampagnen_werden_nicht_zusammengefuehrt(welt, tmp_path):
    ketten = _ketten(welt)
    engine = create_db_engine(welt)
    with engine.begin() as v:  # P steckt in Kampagne X, C1 in Kampagne Y
        for schluessel in (P, C1):
            kid = v.execute(insert(campaigns).values(
                underlying="DAX", eroeffnet_am=datetime(2026, 9, 1, 10, 0))).inserted_primary_key[0]
            v.execute(insert(campaign_links).values(
                campaign_id=kid, chain_id=ketten[schluessel]["chain_id"],
                beziehungstyp="ic_leg", erkennung="manuell", bestaetigt_am=JETZT))
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: "ja"})  # das Paar (P, C1) läge in zwei verschiedenen Kampagnen

    with pytest.raises(ValueError, match="Kampagnen"):
        kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)

    assert len(_alle(welt, campaigns)) == 2  # nichts zusammengeführt, nichts Neues
    assert len(_alle(welt, campaign_links)) == 2


# ---------------------------------------------------------------- Einlesen: Ablehnung


def test_ablehnung_wird_gespeichert_und_verschwindet_aus_der_liste(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: "nein"})

    ergebnis = kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)

    assert (ergebnis.neu_abgelehnt, ergebnis.kampagnen_neu) == (1, 0)
    ketten = _ketten(welt)
    (ablehnung,) = _alle(welt, campaign_ablehnungen)
    assert (ablehnung["chain_id_a"], ablehnung["chain_id_b"], ablehnung["typ"]) == (
        ketten[P]["chain_id"], ketten[C1]["chain_id"], "ic")
    assert ablehnung["erstellt_am"] == JETZT
    assert _alle(welt, campaigns) == []
    # Rebuild zählt den abgelehnten Kandidaten nicht mehr als offen ...
    assert rebuild(welt).campaign_kandidaten_offen == 1
    # ... und die nächste Datei enthält ihn nicht mehr.
    zeilen = _zeilen(_datei(welt, tmp_path, "zweite.xlsx"))
    assert [(z["Typ"], z["Schlüssel Kette B"]) for z in zeilen] == [FIN]


def test_ablehnung_gilt_nur_fuer_den_typ_des_paares(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: "nein"})
    kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)

    zeilen = _zeilen(_datei(welt, tmp_path, "zweite.xlsx"))

    assert len(zeilen) == 1 and zeilen[0]["Typ"] == "financing"


def test_bestaetigung_gegen_gespeicherte_ablehnung_braucht_ueberschreiben(welt, tmp_path):
    erste = _datei(welt, tmp_path)
    _entscheide(erste, {IC: "nein"})
    kampagnen_urteile_einlesen(welt, erste, jetzt=JETZT)
    # Die Datei mit der abgelehnten Zeile erneut ausschreiben ist nicht möglich
    # (sie fehlt); der Widerspruch entsteht durch Wiederverwenden der ersten.
    mappe = load_workbook(erste)
    blatt = mappe["Kandidaten"]
    kopf = [c.value for c in blatt[1]]
    for zeile in range(2, blatt.max_row + 1):
        if blatt.cell(zeile, kopf.index("Typ") + 1).value == "ic":
            blatt.cell(zeile, kopf.index("Entscheidung") + 1, "ja")
    zweite = tmp_path / "widerspruch.xlsx"
    mappe.save(zweite)

    with pytest.raises(ValueError, match="widerspricht"):
        kampagnen_urteile_einlesen(welt, zweite, jetzt=JETZT)
    assert len(_alle(welt, campaign_ablehnungen)) == 1
    assert _alle(welt, campaigns) == []

    ergebnis = kampagnen_urteile_einlesen(welt, zweite, ueberschreiben=True, jetzt=JETZT)

    assert (ergebnis.geaendert, ergebnis.kampagnen_neu) == (1, 1)
    assert _alle(welt, campaign_ablehnungen) == []
    assert len(_alle(welt, campaigns)) == 1


def test_ablehnung_eines_paares_in_derselben_kampagne_bricht_ab(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: "ja"})
    kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)
    _entscheide(ziel, {IC: "nein"})

    with pytest.raises(ValueError, match="Kampagne"):
        kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)

    assert _alle(welt, campaign_ablehnungen) == []


# ---------------------------------------------------------------- Einlesen: Prüfungen


@pytest.mark.parametrize("wert", ["vielleicht", True, 1])
def test_ungueltige_entscheidung_bricht_ab(welt, tmp_path, wert):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: wert})

    with pytest.raises(ValueError, match="Entscheidung"):
        kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)

    assert _alle(welt, campaigns) == []


def test_ungueltiger_typ_bei_bestaetigung_bricht_ab(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: "ja"}, **{"Typ Kette B": {IC: "gegenfinanzierung"}})

    with pytest.raises(ValueError, match="Typ"):
        kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)

    assert _alle(welt, campaigns) == []


def test_unbekannter_schluessel_bricht_ab(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: "ja"}, **{"Schlüssel Kette B": {IC: "k:gibt|es|nicht"}})

    with pytest.raises(ValueError, match="Schlüssel"):
        kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)


def test_leere_entscheidung_schreibt_nichts(welt, tmp_path):
    ziel = _datei(welt, tmp_path)

    ergebnis = kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)

    assert ergebnis.unentschieden == 2
    assert _alle(welt, campaigns) == []
    assert _alle(welt, campaign_ablehnungen) == []


# ---------------------------------------------------------------- Rebuild


def test_klassifikation_und_ketten_ids_ueberleben_den_rebuild(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: "ja"})
    kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)
    vorher = (_alle(welt, campaign_links), _ketten(welt)[P]["chain_id"])

    ergebnis = rebuild(welt)

    assert (_alle(welt, campaign_links), _ketten(welt)[P]["chain_id"]) == vorher
    # IC entschieden (beide Ketten in einer Kampagne), Financing (P, C2) offen.
    assert ergebnis.campaign_kandidaten_offen == 1


def test_rebuild_zieht_die_zeiten_der_kampagne_nach(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: "ja"})
    kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)
    engine = create_db_engine(welt)
    with engine.begin() as v:  # Zeiten verfälschen
        v.execute(update(campaigns).values(
            eroeffnet_am=datetime(2020, 1, 1), beendet_am=datetime(2020, 2, 1)))

    rebuild(welt)

    (kampagne,) = _alle(welt, campaigns)
    assert kampagne["eroeffnet_am"] == datetime(2026, 9, 1, 10, 0)
    assert kampagne["beendet_am"] is None


def test_rebuild_loescht_keine_kette_mit_klassifikation(welt, tmp_path):
    ziel = _datei(welt, tmp_path)
    _entscheide(ziel, {IC: "ja", FIN: "nein"})
    kampagnen_urteile_einlesen(welt, ziel, jetzt=JETZT)
    engine = create_db_engine(welt)
    with engine.begin() as v:  # Die Ausführungen von C1 und C2 verschwinden
        ids = ["F1", "F2", "G1", "G2"]
        v.execute(delete(legs).where(legs.c.open_exec_id.in_(ids)))
        v.execute(delete(executions).where(executions.c.exec_id.in_(ids)))

    rebuild(welt)  # darf nicht an Fremdschlüsseln scheitern

    ketten = _ketten(welt)
    assert C1 in ketten and C2 in ketten
    assert len(_alle(welt, campaign_links)) == 2
    assert len(_alle(welt, campaign_ablehnungen)) == 1


# ---------------------------------------------------------------- CLI


def test_cli_rundreise(welt, tmp_path, capsys):
    ziel = tmp_path / "k.xlsx"

    assert main(["kampagnen-ausschreiben", str(ziel), "--db", str(welt)]) == 0
    assert "1 Gruppen, 2 Kandidaten" in capsys.readouterr().out
    _entscheide(ziel, {IC: "ja"})
    assert main(["kampagnen-einlesen", str(ziel), "--db", str(welt)]) == 0
    ausgabe = capsys.readouterr().out
    assert "1 neu bestätigt" in ausgabe
    assert len(_alle(welt, campaigns)) == 1


def test_cli_meldet_fehler_ohne_traceback(welt, tmp_path, capsys):
    ziel = tmp_path / "k.xlsx"
    main(["kampagnen-ausschreiben", str(ziel), "--db", str(welt)])
    capsys.readouterr()

    assert main(["kampagnen-ausschreiben", str(ziel), "--db", str(welt)]) == 1
    assert "Ausschreiben fehlgeschlagen" in capsys.readouterr().err
    _entscheide(ziel, {IC: "vielleicht"})
    assert main(["kampagnen-einlesen", str(ziel), "--db", str(welt)]) == 1
    assert "Einlesen fehlgeschlagen" in capsys.readouterr().err
