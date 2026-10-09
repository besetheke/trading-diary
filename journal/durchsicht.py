"""Durchsicht der Rollkandidaten über eine Datei-Rundreise (AP-04d).

Seit AP-04k verketten Rolls ohne gemeinsame Order-Nummer nicht mehr
automatisch; sie sind Kandidaten, über die der Nutzer entscheidet (Plan §6,
Block B, Regel 2). Dieses Modul gibt den Weg über hunderte Kandidaten:
``kandidaten_ausschreiben()`` schreibt die offenen als XLSX-Tabelle,
gruppiert nach schließender Position mit allen Fortsetzungen; der Nutzer
entscheidet in der Tabellenkalkulation; ``urteile_einlesen()`` schreibt die
Entscheidungen nach ``chain_overrides``.

Verbindlich ist der Vertrag in ``auftraege/AP-04dv.md`` (Nummern in den
Kommentaren folgen ihm). Das Modul steht bewusst nicht unter ``engine/`` —
es liest und schreibt Dateien und eine Datenbank (Regel 6). Kandidaten werden
nicht neu erfunden: Lesen und Rechnen kommen aus ``bestand_berechnen()`` in
``journal.rebuild`` — ein Lesepfad, kein Abkömmling (Vertrag 10).

Dokumentierte Annahmen (Regel 10 — das Abnahmekriterium zur ``git diff``
lässt neben den drei Modulen keine Änderung an ``OFFENE_FRAGEN.md`` zu):

- **Die Nettoprämie steht in ``journal.preise``** (AP-07p, P5): Seitdem
  rechnet dieses Modul sie nicht mehr selbst, sondern ruft ``nettopraemie``.
  Gerundet wird dort ausdrücklich auf acht Nachkommastellen (P3 in
  ``auftraege/AP-07pv.md``), nicht mehr auf den Cent — ein stiller Schnitt
  auf 28 signifikante Stellen wäre eine Rundung ohne Namen (Regel 5).
- **Dieselbe Zeile zweimal in einer Datei** wird wie eine sequenzielle
  Folge behandelt: gleiche Entscheidung zählt unverändert, eine andere
  folgt der Widerspruchs-Regel (Vertrag 14). Der Vertrag schweigt dazu;
  Ausschreiben erzeugt solche Zeilen nie.
"""

import pathlib
from dataclasses import dataclass
from datetime import datetime

from openpyxl import Workbook, load_workbook
from sqlalchemy import insert, select, update

from engine.chains import _seite_bestimmen
from journal.db import STANDARD_BUSY_TIMEOUT_MS, create_db_engine
from journal.preise import nettopraemie
from journal.rebuild import (
    _kette_schluessel,
    _struktur_schluessel,
    bestand_berechnen,
)
from journal.schema import chain_overrides, chains, structures

# Vertrag 8, Regel 2: Aus der echten Datenbank enthält die Datei echte
# Kontodaten — sie darf nicht im Repository landen (nur unter
# fixtures/golden/, das nicht versioniert wird).
PROJEKTWURZEL = pathlib.Path(__file__).resolve().parents[1]
GOLDEN_VERZEICHNIS = PROJEKTWURZEL / "fixtures" / "golden"

# Vertrag 3: feste Spaltenreihenfolge, Kopfzeile in Zeile 1.
KOPF = [
    "Gruppe", "Kandidaten in Gruppe", "Rang", "Basiswert", "Seite",
    "Alt: Legs", "Alt: geschlossen", "Alt: Preis zu",
    "Neu: Legs", "Neu: eröffnet", "Neu: Preis auf",
    "Abstand Minuten", "Mengenverhältnis", "Entscheidung", "Bemerkung",
    "Schlüssel Kette", "Schlüssel Struktur",
]

# Vertrag 12: umgebende Leerzeichen und Groß-/Kleinschreibung egal; alles
# andere — auch Wahrheitswerte und Zahlen — ist Abbruch.
BESTAETIGT_TEXT = frozenset({"ja", "j"})
ABGELEHNT_TEXT = frozenset({"nein", "n"})


@dataclass(frozen=True)
class AusschreibErgebnis:
    """Zählwerte des Ausschreibens (Vertrag 1)."""

    gruppen: int
    kandidaten: int


@dataclass(frozen=True)
class EinleseErgebnis:
    """Zählwerte des Einlesens (Vertrag 1)."""

    neu_bestaetigt: int
    neu_abgelehnt: int
    unveraendert: int
    geaendert: int
    unentschieden: int


# ---------------------------------------------------------------------------
# Ausschreiben (Vertrag 2–9)
# ---------------------------------------------------------------------------


def kandidaten_ausschreiben(
    db_pfad: str | pathlib.Path,
    ziel: str | pathlib.Path,
    *,
    busy_timeout_ms: int = STANDARD_BUSY_TIMEOUT_MS,
) -> AusschreibErgebnis:
    """Schreibt die offenen Rollkandidaten als XLSX nach ``ziel`` (Vertrag 3).

    Vor jedem Schreibversuch: Existiert das Ziel, ``FileExistsError``
    (Vertrag 7); liegt es im Projektverzeichnis, aber nicht unter
    ``fixtures/golden/``, ``ValueError`` (Vertrag 8). Fehlen die Zeilen zu
    den Naturschlüsseln in ``structures``/``chains``, ``ValueError`` mit dem
    Hinweis auf ``journal rebuild`` (Vertrag 9) — dann entstünde keine Datei.
    """
    ziel = pathlib.Path(ziel)
    if ziel.exists():
        raise FileExistsError(
            f"Zieldatei {ziel} existiert bereits — eine vorhandene Datei kann "
            "Entscheidungen tragen (Vertrag 7)"
        )
    aufgeloest = ziel.resolve()
    if aufgeloest.is_relative_to(PROJEKTWURZEL) and not aufgeloest.is_relative_to(
        GOLDEN_VERZEICHNIS
    ):
        raise ValueError(
            f"Ziel {ziel} liegt im Projektverzeichnis — eine Kandidatendatei "
            "aus der echten Datenbank enthält echte Kontodaten und gehört "
            "nicht ins Repository (Regel 2, Vertrag 8)"
        )

    engine = create_db_engine(db_pfad, busy_timeout_ms=busy_timeout_ms)
    try:
        with engine.connect() as verbindung:
            gruppen = _gruppen_bilden(verbindung)
    finally:
        engine.dispose()

    _mappe_schreiben(ziel, gruppen)
    return AusschreibErgebnis(
        gruppen=len(gruppen), kandidaten=sum(len(zeilen) for zeilen in gruppen.values())
    )


def _gruppen_bilden(verbindung) -> dict[str, list[dict]]:
    """Offene Kandidaten, gruppiert nach schließender Struktur (Vertrag 2/4).

    Die Kandidaten kommen ungefiltert aus der Engine (``bestand_berechnen``,
    Vertrag 10) und sind dort bereits nach (Schließzeitpunkt, Rang) sortiert —
    die Nummerierung der Gruppen ergibt sich aus der Reihenfolge des ersten
    Auftretens. Dieselbe ``entschieden``-Regel wie ``_offene_kandidaten()``.
    """
    bestand = bestand_berechnen(verbindung)
    overrides = [
        dict(z) for z in verbindung.execute(select(chain_overrides)).mappings()
    ]
    struktur_id_je_schluessel = {
        z["natur_schluessel"]: z["structure_id"]
        for z in verbindung.execute(
            select(structures.c.structure_id, structures.c.natur_schluessel).where(
                structures.c.sicht == "konto"
            )
        ).mappings()
    }
    ketten_id_je_schluessel = {
        z["natur_schluessel"]: z["chain_id"]
        for z in verbindung.execute(
            select(chains.c.chain_id, chains.c.natur_schluessel).where(
                chains.c.sicht == "konto"
            )
        ).mappings()
    }
    kette_je_struktur = {
        _struktur_schluessel(glied.structure): _kette_schluessel(kette)
        for kette in bestand.ketten_ergebnis.chains
        for glied in kette.glieder
    }
    leg_je_schluessel = {
        (leg.open_exec_id, leg.close_exec_id): leg for leg in bestand.matching.legs
    }
    entschieden = {(o["chain_id"], o["structure_id"]) for o in overrides}

    gruppen: dict[str, list[dict]] = {}
    for kandidat in bestand.ketten_ergebnis.kandidaten:
        schliessende = _struktur_schluessel(kandidat.schliessende)
        eroeffnende = _struktur_schluessel(kandidat.eroeffnende)
        ketten_schluessel = kette_je_struktur.get(schliessende)
        chain_id = ketten_id_je_schluessel.get(ketten_schluessel)
        structure_id = struktur_id_je_schluessel.get(eroeffnende)
        # Vertrag 9: Die Schlüssel müssen auf vorhandene Zeilen zeigen, sonst
        # ließe sich die Datei nie zurücklesen.
        if chain_id is None or structure_id is None:
            raise ValueError(
                f"Kandidat {schliessende} → {eroeffnende} verweist auf keine "
                "vorhandene Zeile in structures/chains — zuerst `journal "
                "rebuild` ausführen (Vertrag 9)"
            )
        if (chain_id, structure_id) in entschieden:
            continue
        gruppen.setdefault(schliessende, []).append(
            _zeile_bilden(kandidat, ketten_schluessel, eroeffnende, bestand, leg_je_schluessel)
        )
    return gruppen


def _zeile_bilden(
    kandidat, ketten_schluessel: str, eroeffnende_schluessel: str, bestand, leg_je_schluessel
) -> dict:
    """Spaltenwerte eines Kandidaten (Vertrag 5/6); Preise als ``Decimal``."""
    schliessende, eroeffnende = kandidat.schliessende, kandidat.eroeffnende
    return {
        "Kandidaten in Gruppe": None,  # setzt _mappe_schreiben je Gruppe
        "Rang": kandidat.rang,
        "Basiswert": schliessende.underlying,
        "Seite": _seite_bestimmen([schliessende]),
        "Alt: Legs": _legs_als_text(schliessende, bestand, leg_je_schluessel),
        "Alt: geschlossen": schliessende.geschlossen_am,
        "Alt: Preis zu": nettopraemie(schliessende, schliessung=True, bestand=bestand,
                                      leg_je_schluessel=leg_je_schluessel),
        "Neu: Legs": _legs_als_text(eroeffnende, bestand, leg_je_schluessel),
        "Neu: eröffnet": eroeffnende.eroeffnet_am,
        "Neu: Preis auf": nettopraemie(eroeffnende, schliessung=False, bestand=bestand,
                                       leg_je_schluessel=leg_je_schluessel),
        "Abstand Minuten": round(kandidat.abstand.total_seconds() / 60),
        "Mengenverhältnis": kandidat.mengenverhaeltnis,
        "Entscheidung": None,
        "Bemerkung": None,
        "Schlüssel Kette": ketten_schluessel,
        "Schlüssel Struktur": eroeffnende_schluessel,
    }


def _legs_als_text(struktur, bestand, leg_je_schluessel) -> str:
    """Lesbarer Text je Leg: Richtung, Menge, Recht, Strike, Verfall
    (Vertrag 5). Richtung und Menge aus der eröffnenden Ausführung bzw. der
    Leg — so steht die Position darin, wie sie aufgebaut wurde."""
    teile = []
    for schluessel in struktur.leg_schluessel:
        leg = leg_je_schluessel[(schluessel.open_exec_id, schluessel.close_exec_id)]
        oeffnung = bestand.exec_je_id[schluessel.open_exec_id]
        instrument = bestand.instrument_je_conid[leg.conid]
        teile.append(
            f"{oeffnung['side']} {leg.menge} {instrument.get('right') or '?'} "
            f"{instrument.get('strike') if instrument.get('strike') is not None else '?'} "
            f"{instrument.get('expiry') or '?'}"
        )
    return "; ".join(teile)


def _mappe_schreiben(ziel: pathlib.Path, gruppen: dict[str, list[dict]]) -> None:
    """XLSX mit den Blättern ``Kandidaten`` und ``Anleitung`` (Vertrag 3)."""
    mappe = Workbook()
    blatt = mappe.active
    blatt.title = "Kandidaten"
    blatt.append(KOPF)
    for nummer, zeilen in enumerate(gruppen.values(), start=1):
        for zeile in zeilen:
            zeile["Gruppe"] = nummer
            zeile["Kandidaten in Gruppe"] = len(zeilen)
            blatt.append([zeile[name] for name in KOPF])

    anleitung = mappe.create_sheet("Anleitung")
    anleitung["A1"] = (
        "Entscheidung: `ja` (oder `j`) bestätigt den Roll, `nein` (oder `n`) "
        "lehnt ihn ab. Eine Zeile leer lassen heißt unentschieden."
    )
    anleitung["A2"] = (
        "Nicht „Wahr“ oder „Falsch“ eintragen — Excel macht daraus "
        "Wahrheitswerte, die das Einlesen ablehnt."
    )
    anleitung["A3"] = (
        "Eine Gruppe ist eine schließende Position mit allen möglichen "
        "Fortsetzungen. Bei „Kandidaten in Gruppe“ größer 1 ist die Lage "
        "mehrdeutig — dort beginnt die Durchsicht."
    )
    anleitung["A4"] = (
        "Nach dem Einlesen (`journal kandidaten-einlesen`) ist "
        "`journal rebuild` auszuführen."
    )
    mappe.save(ziel)


# ---------------------------------------------------------------------------
# Einlesen (Vertrag 11–17)
# ---------------------------------------------------------------------------


def urteile_einlesen(
    db_pfad: str | pathlib.Path,
    quelle: str | pathlib.Path,
    *,
    ueberschreiben: bool = False,
    jetzt: datetime | None = None,
    busy_timeout_ms: int = STANDARD_BUSY_TIMEOUT_MS,
) -> EinleseErgebnis:
    """Liest Entscheidungen aus einer ausgeschriebenen Datei ein und schreibt
    sie nach ``chain_overrides`` (Vertrag 11).

    Erst werden alle Zeilen geprüft, dann in einer kurzen Transaktion
    geschrieben (Vertrag 15, Regel 9): Bricht irgendetwas ab — ungültiger
    Wert (12), Widerspruch ohne ``ueberschreiben`` (14), doppelte
    Strukturbeanspruchung (16), unbekannter Schlüssel (17) —, steht danach
    nichts Neues in der Datenbank. Kein ``rebuild``; der ist Sache des
    Aufrufenden.
    """
    if jetzt is None:
        jetzt = datetime.now()

    zeilen = _entscheidungen_lesen(pathlib.Path(quelle))
    engine = create_db_engine(db_pfad, busy_timeout_ms=busy_timeout_ms)
    try:
        with engine.begin() as verbindung:
            return _urteile_schreiben(verbindung, zeilen, ueberschreiben=ueberschreiben,
                                      jetzt=jetzt)
    finally:
        engine.dispose()


def _entscheidungen_lesen(quelle: pathlib.Path) -> list[dict]:
    """Zeilen des Blatts ``Kandidaten`` als Abbild der Spalten (Vertrag 3)."""
    mappe = load_workbook(quelle, data_only=True)
    if "Kandidaten" not in mappe.sheetnames:
        raise ValueError(f"Blatt 'Kandidaten' fehlt in {quelle} — keine Entscheidungen")
    blatt = mappe["Kandidaten"]
    kopf = [zelle.value for zelle in blatt[1]]
    fehlend = [name for name in KOPF if name not in kopf]
    if fehlend:
        raise ValueError(
            f"Kopfzeile unvollständig (fehlt: {', '.join(fehlend)}) — keine "
            "Kandidatendatei"
        )
    spalte_je_name = {name: kopf.index(name) for name in KOPF}

    zeilen = []
    for werte in blatt.iter_rows(min_row=2, values_only=True):
        if not any(wert is not None for wert in werte):
            continue
        zeilen.append({name: werte[index] for name, index in spalte_je_name.items()})
    return zeilen


def _urteile_schreiben(verbindung, zeilen: list[dict], *, ueberschreiben: bool,
                       jetzt: datetime) -> EinleseErgebnis:
    """Prüft alles (Vertrag 12/14/16/17) und schreibt erst dann (Vertrag 15)."""
    struktur_id_je_schluessel = {
        z["natur_schluessel"]: z["structure_id"]
        for z in verbindung.execute(
            select(structures.c.structure_id, structures.c.natur_schluessel).where(
                structures.c.sicht == "konto"
            )
        ).mappings()
    }
    ketten_id_je_schluessel = {
        z["natur_schluessel"]: z["chain_id"]
        for z in verbindung.execute(
            select(chains.c.chain_id, chains.c.natur_schluessel).where(
                chains.c.sicht == "konto"
            )
        ).mappings()
    }
    schluessel_je_ketten_id = {cid: s for s, cid in ketten_id_je_schluessel.items()}
    schluessel_je_struktur_id = {sid: s for s, sid in struktur_id_je_schluessel.items()}
    vorhandene = {
        (o["chain_id"], o["structure_id"]): o
        for o in verbindung.execute(select(chain_overrides)).mappings()
    }

    neu_bestaetigt = neu_abgelehnt = unveraendert = geaendert = unentschieden = 0
    einzufuegen: dict[tuple[int, int], tuple[str, str | None]] = {}
    zu_aendern: dict[tuple[int, int], tuple[str, str | None]] = {}
    zustand = {schluessel: zeile["urteil"] for schluessel, zeile in vorhandene.items()}

    for zeile in zeilen:
        # Vertrag 17: unbekannte Schlüssel brechen ab.
        chain_id = ketten_id_je_schluessel.get(zeile["Schlüssel Kette"])
        structure_id = struktur_id_je_schluessel.get(zeile["Schlüssel Struktur"])
        if chain_id is None:
            raise ValueError(
                f"Gruppe {zeile['Gruppe']}, Rang {zeile['Rang']}: unbekannter "
                f"Kettenschlüssel {zeile['Schlüssel Kette']!r} (Vertrag 17)"
            )
        if structure_id is None:
            raise ValueError(
                f"Gruppe {zeile['Gruppe']}, Rang {zeile['Rang']}: unbekannter "
                f"Strukturschlüssel {zeile['Schlüssel Struktur']!r} (Vertrag 17)"
            )

        urteil = _urteil_aus_wert(zeile["Entscheidung"], zeile["Gruppe"], zeile["Rang"])
        if urteil is None:
            # Vertrag 11: leer lassen heißt unentschieden — keine Zeile.
            unentschieden += 1
            continue
        bemerkung = zeile["Bemerkung"]
        notiz = bemerkung if isinstance(bemerkung, str) and bemerkung.strip() else None

        schluessel = (chain_id, structure_id)
        vorher = zustand.get(schluessel)
        if vorher is None:
            einzufuegen[schluessel] = (urteil, notiz)
            if urteil == "bestaetigt":
                neu_bestaetigt += 1
            else:
                neu_abgelehnt += 1
        elif vorher == urteil:
            # Vertrag 13: dieselbe Entscheidung ist kein Fehler, keine neue
            # Zeile.
            unveraendert += 1
        elif not ueberschreiben:
            # Vertrag 14, Regel 4: keine stille Änderung einer getroffenen
            # Entscheidung.
            raise ValueError(
                f"Gruppe {zeile['Gruppe']}, Rang {zeile['Rang']}: Entscheidung "
                f"{'bestätigt' if urteil == 'bestaetigt' else 'abgelehnt'} "
                f"widerspricht der gespeicherten ({vorher}) — erneut einlesen "
                "nur mit ueberschreiben=True"
            )
        elif schluessel in einzufuegen:
            # Annahme aus dem Modul-Docstring: dieselbe Zeile zweimal in
            # einer Datei — die spätere gewinnt, solange noch nichts steht.
            einzufuegen[schluessel] = (urteil, notiz)
            geaendert += 1
        else:
            zu_aendern[schluessel] = (urteil, notiz)
            geaendert += 1
        zustand[schluessel] = urteil

    # Vertrag 16 (AP-04v, Vertrag 21): Keine Struktur in zwei Ketten — aus
    # der Datei selbst oder zusammen mit bereits gespeicherten Bestätigungen.
    endzustand = {k: v["urteil"] for k, v in vorhandene.items()}
    endzustand.update({k: v[0] for k, v in einzufuegen.items()})
    endzustand.update({k: v[0] for k, v in zu_aendern.items()})
    ketten_je_struktur: dict[int, set[int]] = {}
    for (chain_id, structure_id), urteil in endzustand.items():
        if urteil == "bestaetigt":
            ketten_je_struktur.setdefault(structure_id, set()).add(chain_id)
    for structure_id, chain_ids in ketten_je_struktur.items():
        if len(chain_ids) > 1:
            namen = ", ".join(sorted(str(schluessel_je_ketten_id.get(cid, cid))
                                     for cid in chain_ids))
            raise ValueError(
                f"Struktur {schluessel_je_struktur_id.get(structure_id, structure_id)} "
                f"würde von zwei Ketten beansprucht: {namen} (Vertrag 16)"
            )

    # -- Schreiben (Vertrag 15): alles geprüft, eine kurze Transaktion.
    for (chain_id, structure_id), (urteil, notiz) in einzufuegen.items():
        verbindung.execute(
            insert(chain_overrides).values(
                chain_id=chain_id,
                structure_id=structure_id,
                urteil=urteil,
                notiz=notiz,
                erstellt_am=jetzt,
            )
        )
    for (chain_id, structure_id), (urteil, notiz) in zu_aendern.items():
        # Vertrag 14: Urteil, Notiz und erstellt_am werden neu gesetzt.
        verbindung.execute(
            update(chain_overrides)
            .where(
                chain_overrides.c.chain_id == chain_id,
                chain_overrides.c.structure_id == structure_id,
            )
            .values(urteil=urteil, notiz=notiz, erstellt_am=jetzt)
        )

    return EinleseErgebnis(
        neu_bestaetigt=neu_bestaetigt,
        neu_abgelehnt=neu_abgelehnt,
        unveraendert=unveraendert,
        geaendert=geaendert,
        unentschieden=unentschieden,
    )


def _urteil_aus_wert(wert, gruppe, rang) -> str | None:
    """``bestaetigt``/``abgelehnt`` aus dem Zellenwert, ``None`` für leer
    (Vertrag 11/12). Alles andere — auch Wahrheitswerte, die Excel aus
    „Wahr“/„Falsch“ macht (``bool`` ist in Python eine Unterklasse von
    ``int``, deshalb die String-Prüfung vor allem anderen), und Zahlen —
    bricht ab und nennt Gruppe und Rang."""
    if wert is None:
        return None
    if not isinstance(wert, str):
        raise ValueError(
            f"Gruppe {gruppe}, Rang {rang}: ungültige Entscheidung {wert!r} — "
            "nur 'ja' oder 'nein' eintragen, leer lassen heißt unentschieden "
            "(Vertrag 12)"
        )
    text = wert.strip().lower()
    if not text:
        return None
    if text in BESTAETIGT_TEXT:
        return "bestaetigt"
    if text in ABGELEHNT_TEXT:
        return "abgelehnt"
    raise ValueError(
        f"Gruppe {gruppe}, Rang {rang}: ungültige Entscheidung {wert!r} — "
        "nur 'ja' oder 'nein' eintragen, leer lassen heißt unentschieden "
        "(Vertrag 12)"
    )
