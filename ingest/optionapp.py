"""OptionApp-Trade-Export → Rohdaten (AP-13, Teil 1).

Liest den XLSX-Trade-Export von OptionApp ein (§6.1, Anhang
``auftraege/AP-13-anhang-oa-format.md``) und legt jeden Trade mit seinen Legs
in ``optionapp_trades`` beziehungsweise ``optionapp_legs`` ab — Rohdaten,
nichts Abgeleitetes. Die Verknüpfung mit den Ketten und der Ergebnisabgleich
brauchen AP-04 und sind nicht Teil dieses Moduls.

Grundsätze (Vertrag aus ``auftraege/AP-13p.md``):

- **Kein Raten.** Fehlendes Blatt, abweichende Kopfzeile, defekte Leg-Grammatik,
  fehlende Klammernummer, unbekannte Zahlenformate: jeder Fehler bricht mit
  Nennung des Gefundenen ab. Die Grammatik ist über 856 Legs ausnahmslos
  stabil — eine Abweichung ist ein Defekt der Quelle.
- **Decimal, nie float** (Regel 5). `Open`/`Close` stehen als Zahl in der
  Zelle und werden über ``Decimal(str(wert))`` gelesen; `Profit` ist Text im
  deutschen Format mit Währungszeichen (`-1.234,56 $`) und wird eigens
  zerlegt: Währungszeichen nach ``profit_waehrung``, Tausenderpunkte weg,
  Komma zu Punkt.
- **Roh erhalten.** `Name`, `Strikes`, `Stops` und `Status` bleiben als Rohwert
  stehen; das vollständige Zeilenbild steht in ``roh_json``. `Status` wird
  nie auf bekannte Werte abgebildet.
- **Zeitstempel sind naiv und kein Schlüssel.** Das Format ist ein drittes
  neben TDP und Flex, die Zeitzone ist ungeprüft — die Verknüpfung stiftet
  ausschließlich die Klammernummer am Ende von `Name` (§6.1).
- **Idempotent.** Dieselbe `oa_trade`-Nummer mit identischem Inhalt ist
  derselbe Trade und wird einmal gebucht; mit abweichendem Inhalt bricht der
  Import mit Nennung der Nummer ab — kein Überschreiben (Regel 4).

Wie der Flex-Import gilt: Die Datei wird vollständig geparst und geprüft, bevor
die erste Zeile geschrieben wird.
"""

import hashlib
import json
import pathlib
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import insert, select, update

from journal.db import create_db_engine
from journal.schema import import_runs, optionapp_legs, optionapp_trades

# Elf Spalten in fester Reihenfolge (Anhang). Der Import prüft die Kopfzeile
# und bricht bei Abweichung ab — kein Raten über Spaltenpositionen.
KOPFZEILE = (
    "Entry Time",
    "Strategy",
    "Name",
    "Strikes",
    "Quantity",
    "Open",
    "Close",
    "Stops",
    "Profit",
    "Close Time",
    "Status",
)

# Je Leg: [-]<Menge>x<C|P><Strike> (<JJMMTT>) — Strike mit optionalem
# Dezimalteil, weil Strike-Konventionen halbe Punkte kennen (Anhang).
LEG_MUSTER = re.compile(
    r"^(-?)(\d+)x([CP])(\d+(?:\.\d+)?) \((\d{6})\)$"
)

# Profit im deutschen Format mit Währungszeichen: `-1.234,56 $`, `610,00 $`.
PROFIT_MUSTER = re.compile(r"^\d{1,3}(?:\.\d{3})*(?:,\d+)?$")
PROFIT_WAEHRUNGS_MUSTER = re.compile(r"\s*([^0-9,\-.]+)\s*$")

# Quantity ist ein Verhältnis n/m; die Gleichheit beider Zahlen wird nicht
# vorausgesetzt (Anhang, Eigenheit 1).
MENGEN_MUSTER = re.compile(r"^\d+/\d+$")

# Die Klammernummer steht am Zeilenende; der freie Text davor kann selbst
# Klammern enthalten — deshalb nur die letzte Gruppe auswerten (Anhang).
NAME_NUMMER_MUSTER = re.compile(r"\[(\d+)\]$")

# Felder, über die ein Trade identifiziert ist — für die Dublettenprüfung
# gegen die Datenbank und innerhalb einer Datei (Vertrag 10).
VERGLEICHSFELDER = (
    "entry_zeit",
    "schluss_zeit",
    "strategie",
    "name_roh",
    "strikes_roh",
    "menge_zaehler",
    "menge_nenner",
    "offen_preis",
    "schluss_preis",
    "stop_roh",
    "stop_wert",
    "profit",
    "profit_waehrung",
    "status",
)


@dataclass(frozen=True)
class ImportErgebnis:
    """Bericht über einen Importlauf: gelesen, neu (AP-13, Teil 1)."""

    zeilen_gesamt: int
    zeilen_neu: int

    @property
    def zeilen_uebersprungen(self) -> int:
        return self.zeilen_gesamt - self.zeilen_neu


@dataclass
class _TradeInhalt:
    """Vollständig geparster Trade: Felder plus Legs in Quellenreihenfolge."""

    felder: dict
    legs: list[dict] = field(default_factory=list)

    def ist_gleich(self, andere: "_TradeInhalt") -> bool:
        return (
            all(self.felder[f] == andere.felder[f] for f in VERGLEICHSFELDER)
            and self.legs == andere.legs
        )


@dataclass
class _DateiInhalt:
    """Vollständig geparste und geprüfte Datei — erst dann wird geschrieben."""

    trades: list[_TradeInhalt] = field(default_factory=list)
    zeitraum_von: date | None = None
    zeitraum_bis: date | None = None


# ---------------------------------------------------------------------------
# Elementare Wertumwandlungen — jeder Fehler nennt die Zeile
# ---------------------------------------------------------------------------


def _textwert(zelle, kontext: str, feld: str, pflicht: bool = True) -> str | None:
    """Textzelle, stripspace; None/leer ist bei Pflichtfeldern ein Fehler."""
    wert = "" if zelle is None else str(zelle).strip()
    if not wert:
        if pflicht:
            raise ValueError(f"{kontext}: Pflichtfeld {feld!r} fehlt oder ist leer")
        return None
    return wert


def _zahlzelle(zelle, kontext: str, feld: str) -> Decimal | None:
    """`Open`/`Close` stehen als Zahl in der Zelle (Vertrag 2).

    ``Decimal(str(wert))`` — openpyxl liefert ``float``, und ``str()`` davon
    ist die kürzeste darstellende Form: 11.49 bleibt 11.49, ohne
    Float-Arithmetik (Regel 5). Steht der Wert ausnahmsweise als Text in der
    Zelle, wird er mit Dezimalpunkt gelesen.
    """
    if zelle is None or (isinstance(zelle, str) and not zelle.strip()):
        return None
    rohwert = zelle.strip() if isinstance(zelle, str) else str(zelle)
    try:
        return Decimal(rohwert)
    except InvalidOperation as fehler:
        raise ValueError(f"{kontext}: {feld!r} ist keine Zahl: {rohwert!r}") from fehler


def _profitwert(zelle, kontext: str) -> tuple[Decimal, str | None]:
    """`Profit` als deutschem Text mit Währungszeichen (Vertrag 2).

    `-1.234,56 $`: Währungszeichen abtrennen (nach ``profit_waehrung``),
    Tausenderpunkte entfernen, Komma zu Punkt.
    """
    rohwert = _textwert(zelle, kontext, "Profit")
    waehrung = None
    treffer = PROFIT_WAEHRUNGS_MUSTER.search(rohwert)
    if treffer is not None:
        waehrung = treffer.group(1)
        rohwert = rohwert[: treffer.start()].rstrip()
    negativ = rohwert.startswith("-")
    if negativ:
        rohwert = rohwert[1:]
    if not PROFIT_MUSTER.match(rohwert):
        raise ValueError(
            f"{kontext}: 'Profit' folgt nicht dem deutschen Zahlenformat: {rohwert!r}"
        )
    wert = Decimal(rohwert.replace(".", "").replace(",", "."))
    return (-wert if negativ else wert), waehrung


def _stopwert(zelle, kontext: str) -> tuple[str, Decimal | None]:
    """`Stops`: `-` heißt „kein Stop", `0` ist ein gesetzter Stop (Vertrag 3)."""
    stop_roh = _textwert(zelle, kontext, "Stops")
    if stop_roh == "-":
        return stop_roh, None
    try:
        return stop_roh, Decimal(stop_roh)
    except InvalidOperation as fehler:
        raise ValueError(f"{kontext}: 'Stops' ist weder '-' noch Zahl: {stop_roh!r}") from fehler


def _mengenwert(zelle, kontext: str) -> tuple[Decimal, Decimal]:
    """`Quantity` als Verhältnis n/m, beide Zahlen getrennt (Vertrag 4)."""
    rohwert = _textwert(zelle, kontext, "Quantity")
    if not MENGEN_MUSTER.match(rohwert):
        raise ValueError(f"{kontext}: 'Quantity' ist kein Verhältnis n/m: {rohwert!r}")
    zaehler, nenner = rohwert.split("/")
    return Decimal(zaehler), Decimal(nenner)


def _zeitwert(zelle, kontext: str, feld: str, pflicht: bool = True) -> datetime | None:
    """`JJJJ/MM/TT hh:mm:ss`, naiv gespeichert, wie geliefert (Vertrag 7)."""
    rohwert = _textwert(zelle, kontext, feld, pflicht=pflicht)
    if rohwert is None:
        return None
    try:
        return datetime.strptime(rohwert, "%Y/%m/%d %H:%M:%S")
    except ValueError as fehler:
        raise ValueError(
            f"{kontext}: {feld!r} folgt nicht dem Format JJJJ/MM/TT hh:mm:ss: {rohwert!r}"
        ) from fehler


def _klammernummer(name: str, kontext: str) -> str:
    """Die letzte Klammergruppe in `Name` ist die OA-Trade-Nummer (Vertrag 6)."""
    treffer = NAME_NUMMER_MUSTER.search(name)
    if treffer is None:
        raise ValueError(
            f"{kontext}: Name {name!r} ohne Nummer in eckigen Klammern — die "
            "Klammernummer ist der einzige Verknüpfungsschlüssel (§6.1); ein Trade "
            "ohne Schlüssel darf nicht still durchlaufen."
        )
    return treffer.group(1)


def _legs_zerlegen(strikes: str, kontext: str) -> list[dict]:
    """Zerlegt die Spalte `Strikes` nach der Leg-Grammatik (Vertrag 5).

    Legs getrennt durch `` / ``, je Leg ``[-]<Menge>x<C|P><Strike> (<JJMMTT>)``.
    Führendes Minus ist short. Jede Abweichung bricht mit Nennung der Zeile
    und des Leg-Ausschnitts ab.
    """
    teile = strikes.split(" / ")
    if not 2 <= len(teile) <= 4:
        raise ValueError(
            f"{kontext}: 'Strikes' zerfällt in {len(teile)} Legs, erwartet sind "
            "zwei bis vier (Anhang)"
        )
    legs = []
    for position, ausschnitt in enumerate(teile, start=1):
        treffer = LEG_MUSTER.match(ausschnitt)
        if treffer is None:
            raise ValueError(
                f"{kontext}: Leg-Ausschnitt {ausschnitt!r} folgt nicht der "
                "Leg-Grammatik [-]<Menge>x<C|P><Strike> (<JJMMTT>) — die Grammatik "
                "ist über 856 Legs ausnahmslos stabil, eine Abweichung ist ein "
                "Defekt der Quelle."
            )
        vorzeichen, menge, right, strike, verfall = treffer.groups()
        try:
            expiry = datetime.strptime(verfall, "%y%m%d").date()
        except ValueError as fehler:
            raise ValueError(
                f"{kontext}: Leg-Ausschnitt {ausschnitt!r} trägt kein lesbares "
                f"Verfallsdatum JJMMTT: {verfall!r}"
            ) from fehler
        legs.append(
            {
                "position": position,
                "richtung": -1 if vorzeichen == "-" else 1,
                "menge": Decimal(menge),
                "right": right,
                "strike": Decimal(strike),
                "expiry": expiry,
                "roh": ausschnitt,
            }
        )
    return legs


# ---------------------------------------------------------------------------
# Datei parsen
# ---------------------------------------------------------------------------


def _kontext(zeilen_nummer: int, oa_trade: str | None) -> str:
    kennung = f"Trade {oa_trade}" if oa_trade else "Trade ohne Nummer"
    return f"Zeile {zeilen_nummer} ({kennung})"


def _trade_zeile(zeile: tuple, zeilen_nummer: int) -> _TradeInhalt:
    if len(zeile) < len(KOPFZEILE):
        raise ValueError(
            f"Zeile {zeilen_nummer}: {len(zeile)} Zellen statt {len(KOPFZEILE)} — "
            "die Spaltenzahl ist fix (Anhang)"
        )
    name_roh = _textwert(zeile[2], f"Zeile {zeilen_nummer}", "Name")
    oa_trade = _klammernummer(name_roh, f"Zeile {zeilen_nummer}")
    kontext = _kontext(zeilen_nummer, oa_trade)

    entry_zeit = _zeitwert(zeile[0], kontext, "Entry Time")
    schluss_zeit = _zeitwert(zeile[9], kontext, "Close Time", pflicht=False)
    menge_zaehler, menge_nenner = _mengenwert(zeile[4], kontext)
    stop_roh, stop_wert = _stopwert(zeile[7], kontext)
    profit, profit_waehrung = _profitwert(zeile[8], kontext)
    strikes_roh = _textwert(zeile[3], kontext, "Strikes")
    felder = {
        "oa_trade": oa_trade,
        "entry_zeit": entry_zeit,
        "schluss_zeit": schluss_zeit,
        "strategie": _textwert(zeile[1], kontext, "Strategy"),
        "name_roh": name_roh,
        "strikes_roh": strikes_roh,
        "menge_zaehler": menge_zaehler,
        "menge_nenner": menge_nenner,
        "offen_preis": _zahlzelle(zeile[5], kontext, "Open"),
        "schluss_preis": _zahlzelle(zeile[6], kontext, "Close"),
        "stop_roh": stop_roh,
        "stop_wert": stop_wert,
        "profit": profit,
        "profit_waehrung": profit_waehrung,
        "status": _textwert(zeile[10], kontext, "Status"),
        # Vollständiges Zeilenbild; Zahlen als Text, damit JSON trägt.
        "roh_json": json.dumps(
            [None if zelle is None else str(zelle) for zelle in zeile],
            ensure_ascii=False,
        ),
    }
    return _TradeInhalt(felder=felder, legs=_legs_zerlegen(strikes_roh, kontext))


def _parse_datei(datei: pathlib.Path) -> _DateiInhalt:
    mappe = load_workbook(datei, read_only=True, data_only=True)
    try:
        if "Trades" not in mappe.sheetnames:
            raise ValueError(
                f"{datei}: Arbeitsblatt 'Trades' fehlt — gefunden: {mappe.sheetnames}"
            )
        blatt = mappe["Trades"]
        zeilen = list(blatt.iter_rows(values_only=True))
    finally:
        mappe.close()

    if not zeilen:
        raise ValueError(f"{datei}: Datei ohne Kopfzeile")
    gefunden = ["" if zelle is None else str(zelle).strip() for zelle in zeilen[0]]
    if gefunden != list(KOPFZEILE):
        raise ValueError(
            f"{datei}: Kopfzeile weicht ab — gefunden: {gefunden!r}, erwartet: "
            f"{list(KOPFZEILE)!r}. Kein Raten über Spaltenpositionen."
        )

    inhalt = _DateiInhalt()
    gesehene: dict[str, _TradeInhalt] = {}
    for zeilen_nummer, zeile in enumerate(zeilen[1:], start=2):
        if all(zelle is None for zelle in zeile):
            continue  # Leere Nachlaufzeilen sind kein Trade
        trade = _trade_zeile(zeile, zeilen_nummer)
        oa_trade = trade.felder["oa_trade"]
        bekannter = gesehene.get(oa_trade)
        if bekannter is not None:
            if bekannter.ist_gleich(trade):
                continue  # Dieselbe Nummer, identischer Inhalt: derselbe Trade
            raise ValueError(
                f"{datei}: Nummer {oa_trade} zweimal mit abweichendem Inhalt — "
                "keine stille Entscheidung, kein Überschreiben (Vertrag 10)."
            )
        gesehene[oa_trade] = trade
        inhalt.trades.append(trade)

    zeiten = [t.felder["entry_zeit"] for t in inhalt.trades]
    schliessungen = [t.felder["schluss_zeit"] for t in inhalt.trades if t.felder["schluss_zeit"]]
    if zeiten:
        inhalt.zeitraum_von = min(zeiten).date()
        inhalt.zeitraum_bis = max(zeiten + schliessungen).date()
    return inhalt


# ---------------------------------------------------------------------------
# Schreiben
# ---------------------------------------------------------------------------


def _vorhandener_trade(verbindung, oa_trade: str) -> _TradeInhalt | None:
    """Liest einen bereits gebuchten Trade mit Legs für den Abgleich."""
    zeile = verbindung.execute(
        select(optionapp_trades).where(optionapp_trades.c.oa_trade == oa_trade)
    ).mappings().one_or_none()
    if zeile is None:
        return None
    legs = [
        {
            "position": leg["position"],
            "richtung": leg["richtung"],
            "menge": leg["menge"],
            "right": leg["right"],
            "strike": leg["strike"],
            "expiry": leg["expiry"],
            "roh": leg["roh"],
        }
        for leg in verbindung.execute(
            select(optionapp_legs)
            .where(optionapp_legs.c.optionapp_trade_id == zeile["id"])
            .order_by(optionapp_legs.c.position)
        ).mappings()
    ]
    return _TradeInhalt(felder=dict(zeile), legs=legs)


def _trade_schreiben(verbindung, trade: _TradeInhalt, lauf_id: int) -> int:
    ergebnis = verbindung.execute(
        insert(optionapp_trades).values(**trade.felder, import_run_id=lauf_id)
    )
    trade_id = ergebnis.inserted_primary_key[0]
    verbindung.execute(
        insert(optionapp_legs),
        [{**leg, "optionapp_trade_id": trade_id} for leg in trade.legs],
    )
    return trade_id


def import_optionapp_file(
    db_pfad: str | pathlib.Path, datei_pfad: str | pathlib.Path
) -> ImportErgebnis:
    """Importiert einen OptionApp-Trade-Export idempotent (AP-13, Teil 1).

    Rückgabe: ``ImportErgebnis`` mit ``zeilen_gesamt`` (gelesene Trades) und
    ``zeilen_neu`` (eingefügte). Dieselbe Nummer mit identischem Inhalt wird
    nicht erneut gebucht; abweichender Inhalt bricht ab.
    """
    datei = pathlib.Path(datei_pfad)
    if not datei.is_file():
        raise FileNotFoundError(f"Importdatei nicht gefunden: {datei}")

    datei_hash = hashlib.sha256(datei.read_bytes()).hexdigest()
    # Vollständig parsen und prüfen — erst dann beginnt die Transaktion.
    inhalt = _parse_datei(datei)

    engine = create_db_engine(db_pfad)
    with engine.begin() as verbindung:
        bekannter_lauf = verbindung.execute(
            select(import_runs.c.id).where(
                import_runs.c.quelle == "optionapp", import_runs.c.datei_hash == datei_hash
            )
        ).first()
        if bekannter_lauf is not None:
            # Gleicher Dateiinhalt: nichts zu tun.
            return ImportErgebnis(zeilen_gesamt=len(inhalt.trades), zeilen_neu=0)

        # Ein OA-Lauf trägt keine IB-Ergebnisse; kein Vergleich darf hier
        # welche suchen (Vertrag 11, OF-21).
        lauf = verbindung.execute(
            insert(import_runs).values(
                quelle="optionapp",
                datei_hash=datei_hash,
                importiert_am=datetime.now(),
                zeitraum_von=inhalt.zeitraum_von,
                zeitraum_bis=inhalt.zeitraum_bis,
                qualitaetsstufe="B",
                realized_pnl_verfuegbar=False,
                zeilen_gesamt=len(inhalt.trades),
                zeilen_neu=0,
            )
        )
        lauf_id = lauf.inserted_primary_key[0]

        zeilen_neu = 0
        for trade in inhalt.trades:
            oa_trade = trade.felder["oa_trade"]
            vorhanden = _vorhandener_trade(verbindung, oa_trade)
            if vorhanden is not None:
                if vorhanden.ist_gleich(trade):
                    continue  # Bereits gebucht, identischer Inhalt
                raise ValueError(
                    f"Nummer {oa_trade} bereits mit abweichendem Inhalt gebucht — "
                    "keine stille Entscheidung, kein Überschreiben (Vertrag 10)."
                )
            _trade_schreiben(verbindung, trade, lauf_id)
            zeilen_neu += 1

        verbindung.execute(
            update(import_runs).where(import_runs.c.id == lauf_id).values(zeilen_neu=zeilen_neu)
        )

    return ImportErgebnis(zeilen_gesamt=len(inhalt.trades), zeilen_neu=zeilen_neu)
