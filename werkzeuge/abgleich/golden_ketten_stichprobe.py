"""AP-04: Stichprobe gerollter Ketten ohne externe Referenz (DAX, ESTX50).

Dritte der drei AP-04-Abnahmeprüfungen (Plan §6): Für den Nicht-OA-Teil gibt es
keine unabhängige Quelle. Eine Stichprobe gerollter Ketten wird einmal von Hand
durchgesehen, bestätigt und danach als Fixture eingefroren.

Das Skript erzeugt dafür eine Prüfdatei ``fixtures/golden/Stichprobe_Ketten.xlsx``
mit drei Blöcken:

- **Heuristik** — Ketten mit mindestens einem Glied, das nur aus zeitlicher Nähe
  erschlossen ist (Zeitfenster, ``erkennung = 'heuristik'``). Hier liegt das
  Fehlerrisiko, deshalb das größere Gewicht.
- **Belegt** — Ketten, deren Rolls sämtlich über die Order-Nummer belegt sind.
- **Gegenprobe** — Paare, die **nicht** verkettet wurden, obwohl eine
  Schließung und eine Eröffnung auf demselben Basiswert und derselben Seite
  innerhalb von 30 Minuten liegen. Sie prüfen die andere Fehlerrichtung:
  übersehene Rolls.

Die Auswahl ist deterministisch (Ordnung nach Hash des Naturschlüssels), damit
ein zweiter Lauf dieselbe Stichprobe liefert.

**Ausnahme von der Regel dieses Verzeichnisses:** Dieses Skript schreibt genau
eine Datei — und zwar nach ``fixtures/golden/``, das nicht im Repository liegt
(``.gitignore``). Die Prüfdatei enthält echte Kontodaten, weil nur echte Daten
von Hand beurteilt werden können; sie verlässt ``fixtures/golden/`` nie.
Auf der Konsole erscheinen ausschließlich Zählwerte (Regel 2). Eine vorhandene
Prüfdatei wird nicht überschrieben — sie kann bereits Urteile enthalten.

Aufruf aus dem Projektverzeichnis:
``.venv\\Scripts\\python.exe werkzeuge\\abgleich\\golden_ketten_stichprobe.py``
"""

import hashlib
import sys
from collections import Counter
from datetime import timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import golden_bestandsabgleich_of23 as _bestandsabgleich  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from engine.chains import OPTIONS_ANLAGEKLASSEN, bilde_chains  # noqa: E402
from engine.matching import bilde_legs  # noqa: E402
from engine.structures import ZEITFENSTER, bilde_structures, mengen_vergleichbar  # noqa: E402

GOLDEN = Path(__file__).resolve().parents[2] / "fixtures" / "golden"
ZIEL = GOLDEN / "Stichprobe_Ketten.xlsx"
BASISWERTE = ("DAX", "ESTX50")
UMFANG = {
    ("DAX", "Heuristik"): 8,
    ("ESTX50", "Heuristik"): 4,
    ("DAX", "Belegt"): 6,
    ("ESTX50", "Belegt"): 2,
}
GEGENPROBE_FENSTER = timedelta(minutes=30)
# Geschichtet nach dem Grund, aus dem die Engine nicht verkettet hat — so
# zeigt die Gegenprobe, welche Regel gegebenenfalls Rolls übersieht.
GEGENPROBE_UMFANG = {
    "Menge abweichend": 4,
    "alt schon fortgesetzt": 3,
    "neu schon angehaengt": 1,
    "3-30 Minuten": 2,
    "unerklaert": 5,
}

KOPF = [
    "Nr", "Block", "Basiswert", "Glied", "Verknüpfung", "Erkennung / Grund",
    "Eröffnet", "Geschlossen", "Schluss durch", "Legs", "Menge",
    "Brutto", "Kommission", "Urteil", "Bemerkung", "Schlüssel",
]
ERKLAERUNG = [
    "So prüfst du: Jede Nummer ist eine Kette (bzw. in der Gegenprobe ein Paar).",
    "Trage je Nummer in der ersten Zeile unter 'Urteil' ein: richtig / falsch / unklar.",
    "richtig = so habe ich es gehandelt (Kette) bzw. das war KEIN Roll (Gegenprobe).",
    "falsch = die Engine liegt daneben — bitte unter 'Bemerkung' kurz, was stimmt.",
    "Spalte 'Schlüssel' nicht ändern: Darüber wird die Stichprobe später eingefroren.",
]


def _ordnung(schluessel: str) -> str:
    return hashlib.sha256(schluessel.encode()).hexdigest()


def _struktur_schluessel(structure) -> str:
    return "|".join(sorted(k.open_exec_id for k in structure.leg_schluessel))


def _legs_text(structure, exec_je_id, instruments) -> str:
    teile = Counter()
    for schluessel in structure.leg_schluessel:
        oeffnung = exec_je_id.get(schluessel.open_exec_id)
        if oeffnung is None:
            continue
        inst = instruments.get(oeffnung["conid"], {})
        vorzeichen = "+" if oeffnung["side"] == "BUY" else "-"
        verfall = inst.get("expiry")
        teile[
            (
                f"{vorzeichen} {inst.get('right') or inst.get('asset_class')} "
                f"{inst.get('strike') or ''} "
                f"{verfall.strftime('%d.%m.%y') if verfall else ''}"
            ).replace("  ", " ").strip()
        ] += 1
    return " / ".join(f"{n}× {t}" if n > 1 else t for t, n in sorted(teile.items()))


def _zeile(nr, block, structure, glied, verknuepfung, erkennung, exec_je_id, instruments,
           schluessel):
    return [
        nr, block, structure.underlying, glied, verknuepfung, erkennung,
        structure.eroeffnet_am, structure.geschlossen_am, structure.schluss_ereignistyp,
        _legs_text(structure, exec_je_id, instruments), float(structure.menge),
        float(structure.brutto_pnl) if structure.brutto_pnl is not None else None,
        float(structure.kommission_gesamt), None, None, schluessel,
    ]


def main() -> None:
    if ZIEL.exists():
        print("Prüfdatei existiert bereits und wird nicht überschrieben "
              "(sie kann Urteile enthalten). Zum Neuerzeugen vorher umbenennen.")
        raise SystemExit(1)

    exec_rows, instruments, corporate_actions, _offen = _bestandsabgleich.sammeln()
    massnahmen, _u, _o = _bestandsabgleich.kapitalmassnahmen_ableiten(
        corporate_actions, instruments
    )
    matching = bilde_legs(exec_rows, instruments, massnahmen)
    legs = [
        {
            "conid": leg.conid, "open_exec_id": leg.open_exec_id,
            "close_exec_id": leg.close_exec_id, "menge": leg.menge,
            "offen_ja_nein": leg.offen_ja_nein, "brutto_pnl": leg.brutto_pnl,
            "kommission_gesamt": leg.kommission_gesamt,
        }
        for leg in matching.legs
    ]
    strukturen = bilde_structures(legs, exec_rows, instruments)
    ketten = bilde_chains(strukturen, legs, exec_rows, instruments).chains
    exec_je_id = {r["exec_id"]: r for r in exec_rows}

    # Ketten-Blöcke
    kandidaten: dict[tuple[str, str], list] = {}
    for kette in ketten:
        if kette.underlying not in BASISWERTE or kette.anzahl_rolls < 1:
            continue
        block = (
            "Heuristik"
            if any(g.erkennung == "heuristik" for g in kette.glieder[1:])
            else "Belegt"
        )
        kandidaten.setdefault((kette.underlying, block), []).append(kette)

    auswahl = []
    for (basiswert, block), umfang in UMFANG.items():
        pool = sorted(
            kandidaten.get((basiswert, block), []),
            key=lambda k: _ordnung("k:" + _struktur_schluessel(k.glieder[0].structure)),
        )
        auswahl.extend((block, k) for k in pool[:umfang])

    # Gegenprobe: nicht verkettete Paare innerhalb von 30 Minuten
    kette_je_struktur = {}
    for kette in ketten:
        for glied in kette.glieder:
            kette_je_struktur[_struktur_schluessel(glied.structure)] = id(kette)
    optionen = [
        s for s in strukturen.structures
        if s.underlying in BASISWERTE and s.asset_class in OPTIONS_ANLAGEKLASSEN
    ]
    fortgesetzt = {
        _struktur_schluessel(g.structure) for k in ketten for g in k.glieder[:-1]
    }
    angehaengt = {
        _struktur_schluessel(g.structure) for k in ketten for g in k.glieder[1:]
    }

    def _grund(alt, neu, abstand) -> str:
        if abstand > ZEITFENSTER:
            return "3-30 Minuten"
        if _struktur_schluessel(alt) in fortgesetzt:
            return "alt schon fortgesetzt"
        if not mengen_vergleichbar(neu.menge, alt.menge):
            return "Menge abweichend"
        if _struktur_schluessel(neu) in angehaengt:
            return "neu schon angehaengt"
        return "unerklaert"

    paare = []
    for alt in optionen:
        if alt.status != "geschlossen" or alt.geschlossen_am is None:
            continue
        for neu in optionen:
            if neu is alt or neu.underlying != alt.underlying or neu.rechte != alt.rechte:
                continue
            abstand = neu.eroeffnet_am - alt.geschlossen_am
            if timedelta(0) <= abstand <= GEGENPROBE_FENSTER and (
                kette_je_struktur.get(_struktur_schluessel(alt))
                != kette_je_struktur.get(_struktur_schluessel(neu))
            ):
                paare.append((_grund(alt, neu, abstand), alt, neu))
    paare.sort(key=lambda p: _ordnung(_struktur_schluessel(p[1]) + ">" + _struktur_schluessel(p[2])))
    gegenprobe = []
    for grund, umfang in GEGENPROBE_UMFANG.items():
        gegenprobe.extend([p for p in paare if p[0] == grund][:umfang])

    # Schreiben
    mappe = Workbook()
    blatt = mappe.active
    blatt.title = "Stichprobe"
    for text in ERKLAERUNG:
        blatt.append([text])
    blatt.append([])
    blatt.append(KOPF)
    kopfzeile = blatt.max_row
    for zelle in blatt[kopfzeile]:
        zelle.font = Font(bold=True)
    urteil_fuellung = PatternFill("solid", fgColor="FFF2CC")

    nr = 0
    for block, kette in auswahl:
        nr += 1
        schluessel = "k:" + _struktur_schluessel(kette.glieder[0].structure)
        for glied in kette.glieder:
            blatt.append(_zeile(nr, block, glied.structure, glied.sequenz,
                                glied.verknuepfungstyp, glied.erkennung,
                                exec_je_id, instruments, schluessel))
            if glied.sequenz == 1:
                blatt.cell(blatt.max_row, KOPF.index("Urteil") + 1).fill = urteil_fuellung
    for grund, alt, neu in gegenprobe:
        nr += 1
        schluessel = "p:" + _struktur_schluessel(alt) + ">" + _struktur_schluessel(neu)
        blatt.append(_zeile(nr, "Gegenprobe", alt, "alt", "nicht verkettet", grund,
                            exec_je_id, instruments, schluessel))
        blatt.cell(blatt.max_row, KOPF.index("Urteil") + 1).fill = urteil_fuellung
        blatt.append(_zeile(nr, "Gegenprobe", neu, "neu", "nicht verkettet", "-",
                            exec_je_id, instruments, schluessel))

    for spalte, breite in zip("ABCDEFGHIJKLMNOP", (5, 11, 9, 6, 13, 11, 17, 17, 12, 48,
                                                    7, 10, 11, 12, 30, 12), strict=True):
        blatt.column_dimensions[spalte].width = breite
    mappe.save(ZIEL)

    # Konsole: nur Zählwerte (Regel 2)
    zaehler = Counter(block for block, _ in auswahl)
    print("Grundgesamtheit gerollter Ketten: "
          + ", ".join(f"{b} {a}: {len(v)}" for (b, a), v in sorted(kandidaten.items())))
    print(f"Stichprobe: Heuristik {zaehler['Heuristik']}, Belegt {zaehler['Belegt']}, "
          f"Gegenprobe {len(gegenprobe)} (von {len(paare)} nicht verketteten Paaren "
          f"innerhalb von 30 Minuten)")
    print("Nicht verkettete Paare je Grund: "
          + ", ".join(f"{g}: {n}" for g, n in Counter(p[0] for p in paare).most_common()))
    print("Prüfdatei geschrieben: fixtures/golden/Stichprobe_Ketten.xlsx")


if __name__ == "__main__":
    main()
