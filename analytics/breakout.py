"""Stand der BO Research Suite für die Station (AP-14c) — nur lesend.

Die BO Research Suite ist ein eigenes Forschungsprojekt mit strenger Governance
(versiegelte Validation-/Holdout-Stichproben, je Hypothese ein Discovery-Lauf).
Die Station liest daraus nur:

- ``PROJECT_STATE.md``: Datum des Stands (``Updated JJJJ-MM-TT``) und der
  Abschnitt ``## Objective``;
- das Hypothesenregister, Blatt „Experiment Backlog“ der Excel-Datei;
- die Zusammenfassung der **Discovery**-Ergebnisse ``data/results/hyp*_discovery_v1.json``
  (Status, Sitzungen, Kosten 0,5, Bootstrap, Gates) — nie die Einzelsätze.

Nie gelesen: alles zu Validation oder Holdout (auch wenn es je eine Datei gäbe),
und es wird nie ein Lauf angestoßen. Ein Ergebnis zeigt Forschungsstand, keinen
Handelsvorteil (so auch das Projekt selbst).
"""

import json
import pathlib
import re

BACKLOG_BLATT = "Experiment Backlog"
SPALTEN = ("id", "frage", "universum", "stufe", "prioritaet", "status", "evidenz", "entscheidung")

# Die Ergebnisdateien sind bis zu mehrere MB groß; der Auszug wird je Datei und
# Änderungszeit gemerkt, damit die Seite nicht bei jedem Aufruf alles neu liest.
_CACHE: dict[tuple[str, float], dict | None] = {}


def _text(zelle) -> str | None:
    if zelle is None:
        return None
    wert = str(zelle).strip()
    return wert or None


def _projektstand(ordner: pathlib.Path) -> tuple[str | None, str | None]:
    datei = ordner / "PROJECT_STATE.md"
    if not datei.is_file():
        return None, None
    text = datei.read_text(encoding="utf-8", errors="replace")
    stand = re.search(r"Updated (\d{4}-\d{2}-\d{2})", text)
    ziel = re.search(r"^## Objective\s*\n+(.*?)(?=\n## |\Z)", text, re.S | re.M)
    return (stand.group(1) if stand else None,
            re.sub(r"\s+", " ", ziel.group(1)).strip() if ziel else None)


def _hypothesen(ordner: pathlib.Path) -> list[dict]:
    from openpyxl import load_workbook
    for excel in sorted(ordner.glob("*command_center*.xlsx")):
        wb = load_workbook(excel, read_only=True, data_only=True)
        try:
            if BACKLOG_BLATT not in wb.sheetnames:
                continue
            zeilen = list(wb[BACKLOG_BLATT].iter_rows(values_only=True))
        finally:
            wb.close()
        kopf = next((i for i, z in enumerate(zeilen) if z and _text(z[0]) == "ID"), None)
        if kopf is None:
            continue
        ergebnis = []
        for z in zeilen[kopf + 1:]:
            if not z or _text(z[0]) is None:
                continue
            ergebnis.append({n: _text(z[i]) if i < len(z) else None
                             for i, n in enumerate(SPALTEN)})
        return ergebnis
    return []


def _gates(daten: dict) -> dict[str, bool]:
    """Die Ergebnisdateien änderten ihr Format von Hypothese zu Hypothese:
    ``acceptance_gates`` (je Gate ein Objekt mit ``passed``), ``gate_outcomes``
    oder ``gates`` (je Gate ein Wahrheitswert). Es zählt, was vorhanden ist."""
    gates = {name: bool(g["passed"]) for name, g in (daten.get("acceptance_gates") or {}).items()
             if isinstance(g, dict) and "passed" in g}
    if gates:
        return gates
    for schluessel in ("gate_outcomes", "gates"):
        flach = {name: g for name, g in (daten.get(schluessel) or {}).items()
                 if isinstance(g, bool)}
        if flach:
            return flach
    return {}


def _status(daten: dict) -> str | None:
    if daten.get("promotion_status"):
        return daten["promotion_status"]
    if "all_gates_pass" in daten:
        return "Alle Gates bestanden" if daten["all_gates_pass"] else "Gates nicht bestanden"
    return None


def _auszug(datei: pathlib.Path) -> dict | None:
    schluessel = (str(datei), datei.stat().st_mtime)
    if schluessel in _CACHE:
        return _CACHE[schluessel]
    try:
        d = json.loads(datei.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        _CACHE[schluessel] = None
        return None
    if d.get("sample_role") not in (None, "discovery"):
        _CACHE[schluessel] = None
        return None
    kosten = (d.get("cost_reports") or {}).get("0.5") or {}
    boot = d.get("bootstrap") or {}
    auszug = {
        "id": d.get("hypothesis_id"),
        "status": _status(d),
        "sitzungen": d.get("session_count"),
        "gesamt_punkte_05": kosten.get("total_points"),
        "mittel_punkte_05": kosten.get("mean_points"),
        "bootstrap_p05": boot.get("primary_p05"),
        "gates": _gates(d),
    }
    _CACHE[schluessel] = auszug
    return auszug


def _ergebnisse(ordner: pathlib.Path) -> list[dict]:
    ergebnis = []
    for datei in sorted((ordner / "data" / "results").glob("hyp*_discovery_v1.json")):
        auszug = _auszug(datei)
        if auszug and auszug["id"]:
            ergebnis.append(auszug)
    return sorted(ergebnis, key=lambda e: e["id"])


def breakout_bericht(ordner: str | pathlib.Path | None) -> dict:
    """Auszug für ``GET /breakout``; ohne Ordner ``gefunden = False``."""
    leer = {"gefunden": False, "stand": None, "ziel": None, "hypothesen": [], "ergebnisse": []}
    if ordner is None or not pathlib.Path(ordner).is_dir():
        return leer
    ordner = pathlib.Path(ordner)
    stand, ziel = _projektstand(ordner)
    return {"gefunden": True, "stand": stand, "ziel": ziel,
            "hypothesen": _hypothesen(ordner), "ergebnisse": _ergebnisse(ordner)}
